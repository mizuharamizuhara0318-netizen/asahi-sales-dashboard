import os
import io
import time
import json
import glob
import re
import base64
import warnings
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import plotly.express as px
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

st.set_page_config(page_title="売上管理＆営業分析表", layout="wide")

st.markdown("""
<style>
    .main-header {
        background-color: #2a75d3;
        color: white;
        padding: 8px 15px;
        font-size: 19px;
        font-weight: bold;
        border-radius: 4px;
        margin-bottom: 10px;
    }
    .total-card {
        background-color: #f0f4f8;
        padding: 10px 14px;
        border-radius: 6px;
        border-left: 5px solid #2a75d3;
        margin-bottom: 8px;
    }
    .est-card {
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        background-color: #ffffff;
        margin-bottom: 16px;
        overflow: hidden;
    }
    .est-header {
        background-color: #f8fafc;
        padding: 10px 14px;
        border-bottom: 1px solid #cbd5e1;
        display: flex;
        justify-content: space-between;
        align-items: center;
    }
    .badge-green {
        background-color: #16a34a;
        color: white;
        padding: 2px 6px;
        border-radius: 3px;
        font-size: 11px;
        font-weight: bold;
    }
    .modal-box-header {
        background: linear-gradient(135deg, #1e3a8a, #2563eb);
        color: white;
        padding: 12px 18px;
        border-radius: 6px 6px 0 0;
        font-size: 16px;
        font-weight: bold;
        margin-bottom: 10px;
    }
    .modal-summary-box {
        background-color: #fefce8;
        border: 1px solid #fde047;
        border-left: 6px solid #eab308;
        padding: 10px 16px;
        font-size: 13px;
        margin-bottom: 12px;
        border-radius: 4px;
    }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-header">売上管理＆営業分析表</div>', unsafe_allow_html=True)

BASE_SHARE = r"G:\マイドライブ\売上ダッシュボード"
OTHER_DATA_DIR = os.path.join(BASE_SHARE, "その他データ")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CACHE_FILE = os.path.join(BASE_DIR, "sales_cache.parquet")
SAVED_QUOTES_FILE = os.path.join(BASE_SHARE, "saved_estimates.json")

STAFF_MAP = {
    "1": "小泉", "2": "河端", "3": "安部", "4": "原",
    "5": "西山", "6": "萩原", "7": "田中", "8": "水原",
    "9": "川村", "10": "穴山", "11": "鈴木", "12": "松平"
}

ORDER_CSV_COLUMNS = [
    '受注ID', '顧客ID', '要望等', '注文者お名前(姓)', '注文者お名前(名)', '注文者お名前(フリガナ・姓)', '注文者お名前(フリガナ名)',
    '注文者会社名', '注文者メールアドレス', '注文者メールアドレスタイプ', '注文者電話番号', '注文者FAX', '注文者国', '注文者郵便番号',
    '注文者都道府県', '注文者住所1', '注文者住所2', '注文者性別ID', '注文者性別', '注文者生年月日', '注文者職種ID', '注文者職種',
    '商品合計', '調整額', '調整額備考', 'うち割引額１', 'うち割引額２', 'うち割引額３', 'うちその他調整額', '合計額', '送料（税別）',
    '手数料（税別）', 'クーポンコード', 'クーポン名', 'クーポン利用', 'ポイント利用', '請求額', '消費税額', '加算ポイント', '入金額',
    '配送コスト', 'EC種類', 'EC種類名', 'EC注文番号', 'EC注文番号枝番', 'EC店舗ID', 'EC店舗名', '販売ルートID', '受注ステータス',
    '支払方法ID', '決済ステータス', '配送方法ID', '配送区分', '申込番号種類', '注文日時', '出荷予定日', '出荷日', 'ひとことメモ',
    '伝達事項', '直接ＡＤコード', '初回ＡＤコード', '再計算フラグ', '送付先コード', '送付先お名前(姓)', '送付先お名前(名)',
    '送付先お名前(フリガナ・姓)', '送付先お名前(フリガナ名)', '送付先会社名', '送付先電話番号', '送付先FAX', '送付先国',
    '送付先郵便番号', '送付先都道府県', '送付先住所1', '送付先住所2', 'お届け日', 'お届け時間', 'マーク１', 'マーク２', 'マーク３',
    '出力フラグ(受注伝票)', '出力フラグ(お手紙)', '出力フラグ(納品書)', '出力フラグ(領収書)', '出力フラグ(請求書)', '出力フラグ(EMS)',
    '出力フラグ(送り状)', '出力済(受注伝票)', '出力済(お手紙)', '出力済(納品書)', '出力済(領収書)', '出力済(請求書)', '出力済(EMS)',
    '出力済(送り状)', '出力済(後払い取引登録)', '出力済(後払い出荷登録)', '出力済(配送指示伝票)', '出力済(モール連携出荷)',
    '出力済(モール連携処理済)', 'メール送信済(サンクス)', 'メール送信済(入金確認)', 'メール送信済(発送完了)', 'メール送信済(フォロー)',
    'メール送信済(その他１)', 'メール送信済(その他２)', 'ご購入明細商品コード', 'ご購入明細商品名', 'ご購入明細消費税区分',
    'ご購入明細規格名(横軸選択肢)', 'ご購入明細規格名(縦軸選択肢)', 'ご購入明細選択肢名', 'ご購入明細消費税課税規則', 'ご購入明細送料区分',
    'ご購入明細手数料区分', 'ご購入明細税抜単価', 'ご購入明細税込単価', 'ご購入明細消費税', 'ご購入明細消費税率', 'ご購入明細数量',
    'ご購入明細金額', 'ご購入明細定期区分', 'ご購入明細標準販売価格', 'ご購入明細JANコード', 'ご購入明細ロットNO', '消費税計算区分'
]

# 🔴 ユーザー指定の厳密な12種類のランク名
ALLOWED_EXACT_RANKS = [
    "1.1倍", "1.15倍", "1.2倍", "1.25倍", "1.3倍", "1.35倍",
    "1.4倍", "1.45倍", "1.5倍", "0.9倍", "1.05倍", "1.07倍"
]

def generate_excel_bytes(data_df, title_str):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        data_df.to_excel(writer, index=False, sheet_name=str(title_str)[:30])
    return output.getvalue()

def clean_cust_name(val):
    if not isinstance(val, str) or not val.strip(): return ""
    raw = str(val).strip()
    norm = re.sub(r'[\s ]+', '', raw)
    if "アマゾン" in norm:
        if "セラー" in norm: return "アマゾンセラー"
        elif "ホーム" in norm or "home" in norm.lower(): return "アマゾンベンダー（ホーム）"
        elif "アクセ" in norm or "acc" in norm.lower(): return "アマゾンベンダー（アクセサリー）"
        elif "ベンダー" in norm: return "アマゾンベンダー（ホーム）"
        else: return "アマゾンセラー"
    return norm

def clean_num(val):
    if pd.isna(val): return 0.0
    s = re.sub(r'[^\d.-]', '', str(val))
    try: return float(s) if s else 0.0
    except: return 0.0

def convert_staff_name(val):
    if pd.isna(val): return ""
    v = str(val).strip().replace(".0", "")
    return STAFF_MAP.get(v, v)

def clean_product_name(val):
    if not isinstance(val, str) or not val.strip(): return ""
    if "】" in val: val = val.split("】", 1)[-1]
    return re.sub(r'[\(（]\s*[\d０-９]{7,16}\s*[\)）]', '', val).strip()

def find_file_multidir(filename_pattern):
    search_dirs = [BASE_DIR, OTHER_DATA_DIR, BASE_SHARE]
    for d in search_dirs:
        if os.path.exists(d):
            fl = glob.glob(os.path.join(d, filename_pattern))
            if not fl:
                fl = glob.glob(os.path.join(d, "**", filename_pattern), recursive=True)
            if fl:
                return fl[0], os.path.getmtime(fl[0])
    return None, 0

def get_cost_master_file_info():
    return find_file_multidir("*原価*.xls*")

@st.cache_data(show_spinner=False)
def load_cost_master(file_path, file_mtime):
    cost_dict = {}
    if file_path and os.path.exists(file_path):
        try:
            cdf = pd.read_excel(file_path)
            c_sku = [c for c in cdf.columns if any(k in str(c) for k in ['商品管理', 'SKU', 'コード', '品番', '商品番号'])]
            c_cost = [c for c in cdf.columns if '原価' in str(c)]
            sku_col = c_sku[0] if c_sku else cdf.columns[0]
            cost_col = c_cost[0] if c_cost else (cdf.columns[1] if len(cdf.columns) > 1 else None)
            if cost_col:
                for _, r in cdf.iterrows():
                    sk = str(r[sku_col]).strip()
                    co = clean_num(r[cost_col])
                    if sk: cost_dict[sk] = co
        except Exception: pass
    return cost_dict

def get_member_price_file_info():
    return find_file_multidir("*加盟店価格*.xls*")

@st.cache_data(show_spinner=False)
def load_member_price_master_v2(file_path, file_mtime):
    cust_sku_map = {}
    default_min_map = {}
    if file_path and os.path.exists(file_path):
        try:
            mdf = pd.read_excel(file_path)
            for _, r in mdf.iterrows():
                sku_mgmt = str(r.iloc[3]).strip() if len(r) > 3 else ""
                sku_num = str(r.iloc[4]).strip() if len(r) > 4 else ""
                skus = [s for s in [sku_mgmt, sku_num] if s and s != "nan"]
                if not skus: continue

                j_price = int(clean_num(r.iloc[9])) if len(r) > 9 else 0
                n_min_price = int(clean_num(r.iloc[13])) if len(r) > 13 else 0
                m_comp = str(r.iloc[12]).strip() if len(r) > 12 else ""

                for s in skus:
                    if 50 <= n_min_price < 90000:
                        default_min_map[s] = n_min_price

                    if m_comp and m_comp not in ["nan", "None", "#N/A", ""]:
                        norm_comp = re.sub(r'[\s ]+', '', m_comp)
                        if 50 <= j_price < 90000:
                            cust_sku_map[(m_comp, s)] = j_price
                            cust_sku_map[(norm_comp, s)] = j_price
        except Exception: pass
    return cust_sku_map, default_min_map

def get_rank_price_file_info():
    for pat in ["*ランク価格*.xls*", "*ランク価格*.csv", "*ランク*.xls*", "*ランク*.csv"]:
        p, mt = find_file_multidir(pat)
        if p:
            return p, mt
    return None, 0

@st.cache_data(show_spinner=False)
def load_rank_price_master(file_path, file_mtime):
    if not file_path or not os.path.exists(file_path):
        return pd.DataFrame()
    try:
        if file_path.lower().endswith(".csv"):
            for enc in ["utf-8-sig", "cp932", "shift_jis", "utf-8"]:
                try:
                    return pd.read_csv(file_path, encoding=enc, low_memory=False)
                except Exception:
                    continue
        else:
            return pd.read_excel(file_path)
    except Exception:
        pass
    return pd.DataFrame()

def get_jan_file_info():
    return find_file_multidir("*JAN*.xls*")

@st.cache_data(show_spinner=False)
def load_jan_master_v2(file_path, file_mtime):
    jan_dict = {}
    full_title_dict = {}
    std_price_dict = {}
    if file_path and os.path.exists(file_path):
        try:
            jdf = pd.read_excel(file_path)
            c_sku = [c for c in jdf.columns if any(k in str(c) for k in ['商品番号', '商品コード', 'SKU', '品番'])]
            c_jan = [c for c in jdf.columns if any(k in str(c) for k in ['カタログID', 'JAN', 'JANコード'])]
            c_name = [c for c in jdf.columns if any(k in str(c) for k in ['商品名', '品名'])]
            c_std = [c for c in jdf.columns if any(k in str(c) for k in ['上代', '標準', '定価', '税抜上代'])]
            s_col = c_sku[0] if c_sku else jdf.columns[0]
            j_col = c_jan[0] if c_jan else (jdf.columns[1] if len(jdf.columns) > 1 else None)
            n_col = c_name[0] if c_name else None
            p_col = c_std[0] if c_std else None

            for _, r in jdf.iterrows():
                sk = str(r[s_col]).strip()
                if not sk or sk == "nan": continue
                if j_col and pd.notna(r[j_col]):
                    try: jan_dict[sk] = str(int(float(r[j_col])))
                    except: jan_dict[sk] = str(r[j_col]).replace(".0", "").strip()
                if n_col and pd.notna(r[n_col]):
                    full_title_dict[sk] = str(r[n_col]).strip()
                if p_col and pd.notna(r[p_col]):
                    std_price_dict[sk] = int(clean_num(r[p_col]))
        except Exception: pass
    return jan_dict, full_title_dict, std_price_dict

def get_customer_file_info():
    return find_file_multidir("*注文者会社名*.xls*")

@st.cache_data(show_spinner=False)
def load_customer_master(file_path, file_mtime):
    if file_path and os.path.exists(file_path):
        try: return pd.read_excel(file_path)
        except Exception: pass
    return pd.DataFrame()

def get_hanko_bytes():
    search_dirs = [BASE_DIR, OTHER_DATA_DIR, BASE_SHARE]
    patterns = ["*hanko*.png", "*05171642*.png", "*印鑑*.png", "input_file_5.png"]
    for d in search_dirs:
        if os.path.exists(d):
            for p in patterns:
                fl = glob.glob(os.path.join(d, p))
                if fl:
                    try:
                        with open(fl[0], "rb") as f: return f.read(), fl[0]
                    except: pass
    return None, None

def load_saved_quotes():
    if os.path.exists(SAVED_QUOTES_FILE):
        try:
            with open(SAVED_QUOTES_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception: return []
    return []

def save_quotes_list(q_list):
    try:
        with open(SAVED_QUOTES_FILE, "w", encoding="utf-8") as f:
            json.dump(q_list, f, ensure_ascii=False, indent=2)
    except Exception: pass

def create_smaregi_order_csv_bytes(est_data, raw_title_map, std_price_map):
    rows = []
    pay_map = {"銀行振込": 3, "代金引換": 2, "現行通り": 3, "掛け払い（月末締め翌月末払い）": 3}
    pay_id = pay_map.get(est_data.get("payment", ""), 3)
    now_dt_str = datetime.now().strftime("%Y/%m/%d %H:%M")
    ship_dt_str = datetime.now().strftime("%Y/%m/%d")
    deliv_dt_str = (datetime.now() + timedelta(days=1)).strftime("%Y/%m/%d")

    for item in est_data.get("items", []):
        sku = item.get("商品コード", "")
        qty = int(item.get("数量", 1))
        unit_p = int(item.get("見積単価", 0))
        line_amt = unit_p * qty
        unit_tax = int(round(unit_p * 0.10))
        unit_tax_in = unit_p + unit_tax
        jan = item.get("JANコード", "")
        clean_name = item.get("商品名", "")
        full_prod_name = raw_title_map.get(sku, f"【{sku}】{clean_name}（{jan}）" if jan else f"【{sku}】{clean_name}")
        std_p = std_price_map.get(sku, "")

        r_dict = {col: "" for col in ORDER_CSV_COLUMNS}
        r_dict['顧客ID'] = est_data.get("cust_id", -1)
        r_dict['要望等'] = est_data.get("memo", "")
        r_dict['注文者お名前(姓)'] = est_data.get("cust_sei", "ご担当者")
        r_dict['注文者お名前(名)'] = est_data.get("cust_mei", "様")
        r_dict['注文者お名前(フリガナ・姓)'] = est_data.get("cust_kana_sei", "")
        r_dict['注文者お名前(フリガナ名)'] = est_data.get("cust_kana_mei", "")
        r_dict['注文者会社名'] = est_data.get("cust_comp", "")
        r_dict['注文者メールアドレス'] = est_data.get("cust_email", "")
        r_dict['注文者メールアドレスタイプ'] = 1
        r_dict['注文者電話番号'] = est_data.get("cust_tel", "")
        r_dict['注文者FAX'] = est_data.get("cust_fax", "")
        r_dict['注文者郵便番号'] = est_data.get("cust_post", "")
        r_dict['注文者都道府県'] = est_data.get("cust_pref", "")
        r_dict['注文者住所1'] = est_data.get("cust_addr1", "")
        r_dict['注文者住所2'] = est_data.get("cust_addr2", "")
        r_dict['注文者性別ID'] = 1
        r_dict['注文者性別'] = "男性"
        r_dict['注文者職種'] = est_data.get("cust_job_id", "")
        r_dict['商品合計'] = line_amt
        r_dict['調整額'] = 0
        r_dict['うち割引額１'] = 0
        r_dict['うち割引額２'] = 0
        r_dict['うち割引額３'] = 0
        r_dict['うちその他調整額'] = 0
        r_dict['合計額'] = line_amt
        r_dict['送料（税別）'] = est_data.get("shipping_fee", 0)
        r_dict['手数料（税別）'] = est_data.get("cod_fee", 0)
        r_dict['クーポン名'] = str(est_data.get("rank_name", "")).split(":")[0].strip()
        r_dict['クーポン利用'] = 0
        r_dict['ポイント利用'] = 0
        r_dict['請求額'] = line_amt
        r_dict['消費税額'] = unit_tax
        r_dict['入金額'] = 0
        r_dict['EC種類'] = 16
        r_dict['EC注文番号'] = est_data.get("q_no", "").replace("-", "")
        r_dict['EC店舗ID'] = 1000
        r_dict['EC店舗名'] = "旭電機化成"
        r_dict['販売ルートID'] = 1
        r_dict['受注ステータス'] = 1
        r_dict['支払方法ID'] = pay_id
        r_dict['決済ステータス'] = 0
        r_dict['配送方法ID'] = 1
        r_dict['配送区分'] = 0
        r_dict['申込番号種類'] = 0
        r_dict['注文日時'] = est_data.get("created_at", now_dt_str)
        r_dict['出荷予定日'] = ship_dt_str
        r_dict['ひとことメモ'] = est_data.get("hitokoto", "")
        r_dict['伝達事項'] = est_data.get("memo", "")
        r_dict['再計算フラグ'] = 1
        r_dict['送付先お名前(姓)'] = est_data.get("ship_sei", est_data.get("cust_sei", "ご担当者"))
        r_dict['送付先お名前(名)'] = est_data.get("ship_mei", est_data.get("cust_mei", "様"))
        r_dict['送付先会社名'] = est_data.get("ship_comp", est_data.get("cust_comp", ""))
        r_dict['送付先電話番号'] = est_data.get("ship_tel", est_data.get("cust_tel", ""))
        r_dict['送付先郵便番号'] = est_data.get("ship_post", est_data.get("cust_post", ""))
        r_dict['送付先都道府県'] = est_data.get("ship_pref", est_data.get("cust_pref", ""))
        r_dict['送付先住所1'] = est_data.get("ship_addr1", est_data.get("cust_addr1", ""))
        r_dict['送付先住所2'] = est_data.get("ship_addr2", est_data.get("cust_addr2", ""))
        r_dict['お届け日'] = deliv_dt_str
        r_dict['出力フラグ(受注伝票)'] = 0
        r_dict['出力フラグ(お手紙)'] = 0
        r_dict['出力フラグ(納品書)'] = 1
        r_dict['出力フラグ(領収書)'] = 0
        r_dict['出力フラグ(請求書)'] = 0
        r_dict['出力フラグ(EMS)'] = 0
        r_dict['出力フラグ(送り状)'] = 0
        for flag_col in [
            '出力済(受注伝票)', '出力済(お手紙)', '出力済(納品書)', '出力済(領収書)', '出力済(請求書)',
            '出力済(EMS)', '出力済(送り状)', '出力済(後払い取引登録)', '出力済(後払い出荷登録)',
            '出力済(配送指示伝票)', '出力済(モール連携出荷)', '出力済(モール連携処理済)',
            'メール送信済(サンクス)', 'メール送信済(入金確認)', 'メール送信済(発送完了)',
            'メール送信済(フォロー)', 'メール送信済(その他１)', 'メール送信済(その他２)'
        ]:
            r_dict[flag_col] = 0

        r_dict['ご購入明細商品コード'] = sku
        r_dict['ご購入明細商品名'] = full_prod_name
        r_dict['ご購入明細消費税区分'] = "別"
        r_dict['ご購入明細消費税課税規則'] = 3
        r_dict['ご購入明細送料区分'] = "別"
        r_dict['ご購入明細手数料区分'] = "別"
        r_dict['ご購入明細税抜単価'] = unit_p
        r_dict['ご購入明細税込単価'] = unit_tax_in
        r_dict['ご購入明細消費税'] = unit_tax
        r_dict['ご購入明細消費税率'] = 10
        r_dict['ご購入明細数量'] = qty
        r_dict['ご購入明細金額'] = line_amt
        r_dict['ご購入明細定期区分'] = "商品"
        r_dict['ご購入明細標準販売価格'] = std_p
        r_dict['消費税計算区分'] = "伝票単位(B)"
        rows.append(r_dict)

    out_df = pd.DataFrame(rows, columns=ORDER_CSV_COLUMNS)
    return out_df.to_csv(index=False, encoding="cp932", errors="replace").encode("cp932")

def _read_single_csv(f):
    path_date = None
    ym_match = re.search(r'(\d{4})年?[/\\_](\d{1,2})月?[/\\_]yayoi_(\d{2})(\d{2})', f, re.IGNORECASE)
    if ym_match:
        y, m, md_m, md_d = ym_match.groups()
        path_date = f"{y}{md_m}{md_d}"
    else:
        y_m = re.search(r'(\d{4})年', f)
        d_m = re.search(r'yayoi_(\d{4})', f, re.IGNORECASE)
        if y_m and d_m: path_date = f"{y_m.group(1)}{d_m.group(1)}"

    for enc in ["cp932", "shift_jis", "utf-8"]:
        try:
            d = pd.read_csv(f, encoding=enc, encoding_errors="replace", low_memory=False)
            cols = [str(c).strip() for c in d.columns]
            qty_cols = [c for c in cols if "数量" in c]
            d["_auto_qty"] = d[qty_cols[0]].apply(clean_num) if qty_cols else (
                d.iloc[:, 21].apply(clean_num) if len(cols) >= 22 else 0.0
            )
            price_cols = [c for c in cols if "単価" in c]
            d["_auto_price"] = d[price_cols[0]].apply(clean_num) if price_cols else (
                d.iloc[:, 20].apply(clean_num) if len(cols) >= 21 else 0.0
            )
            amt_cols = [c for c in cols if c in ["金額", "売上金額", "販売金額"]]
            if amt_cols:
                d["_auto_amt"] = d[amt_cols[0]].apply(clean_num)
                d["_auto_amt"] = d.apply(lambda r: (r["_auto_qty"] * r["_auto_price"]) if r["_auto_amt"] == 0 and r["_auto_price"] > 0 else r["_auto_amt"], axis=1)
            else:
                d["_auto_amt"] = d["_auto_qty"] * d["_auto_price"]
            d["_fallback_date"] = path_date
            d["_source_file"] = os.path.basename(f)
            return d
        except Exception: continue
    return None

def _read_amazon_seller_csv(f):
    norm_f = f.replace('\\', '/')
    ym_match = re.search(r'(\d{4})年?[/](\d{1,2})月?', norm_f)
    if ym_match:
        y_str = ym_match.group(1)
        m_str = f"{int(ym_match.group(2)):02d}"
        target_ym = f"{y_str}年{m_str}月"
        try: dt_last = pd.Timestamp(year=int(y_str), month=int(m_str), day=1) + pd.offsets.MonthEnd(1)
        except Exception: dt_last = datetime(int(y_str), int(m_str), 28)
    else:
        target_ym = "2026年09月"
        dt_last = datetime(2026, 9, 30)

    for enc in ["utf-8-sig", "utf-8", "cp932", "shift_jis"]:
        try:
            d = pd.read_csv(f, encoding=enc, encoding_errors="replace", low_memory=False)
            if d.empty: continue
            sku_series = None
            for c in d.columns:
                if str(c).strip().upper() == "SKU":
                    sku_series = d[c].astype(str).str.strip()
                    break
            if sku_series is None:
                sku_series = d.iloc[:, 3].astype(str).str.strip() if len(d.columns) >= 4 else d.iloc[:, 0].astype(str).str.strip()

            qty_series = None
            for c in d.columns:
                if any(k in str(c) for k in ['点数', '数量', '個数']):
                    qty_series = d[c].apply(clean_num).astype(int)
                    break
            if qty_series is None:
                if len(d.columns) >= 15: qty_series = d.iloc[:, 14].apply(clean_num).astype(int)
                elif len(d.columns) >= 3: qty_series = d.iloc[:, 2].apply(clean_num).astype(int)
                else: qty_series = pd.Series([1] * len(d))

            amt_series = None
            for c in d.columns:
                if any(k in str(c) for k in ['売上額', '売上金額', '売上', '金額']):
                    gross = d[c].apply(clean_num)
                    amt_series = (gross / 1.10).round().astype(int)
                    break
            if amt_series is None:
                if len(d.columns) >= 19: gross = d.iloc[:, 18].apply(clean_num); amt_series = (gross / 1.10).round().astype(int)
                elif len(d.columns) >= 2: gross = d.iloc[:, 1].apply(clean_num); amt_series = (gross / 1.10).round().astype(int)
                else: continue

            res = pd.DataFrame({
                "商品コード": sku_series, "商品名": sku_series, "得意先名": "アマゾンセラー", "担当者": "水原",
                "数量": qty_series, "金額": amt_series, "日付_dt": dt_last,
                "日付": dt_last.strftime('%Y/%m/%d'), "年月": target_ym,
                "_source_file": os.path.basename(f)
            })
            res["単価"] = (res["金額"] / res["数量"].replace(0, 1)).round().astype(int)
            return res
        except Exception: continue
    return None

def _parse_raw_df_to_standard(df_all):
    if df_all.empty: return pd.DataFrame()
    def find_col(candidates, cols):
        for cand in candidates:
            for c in cols:
                if cand == c or cand in c: return c
        return None

    all_cols = [str(c).strip() for c in df_all.columns]
    c_date = find_col(["納品日", "伝票日付", "売上日", "日付"], all_cols)
    c_pcode = find_col(["商品コード", "商品CD", "品番"], all_cols)
    c_pname = find_col(["商品名/入金内容", "商品名", "商品名_摘要", "品名"], all_cols)
    c_cname = find_col(["得意先名称", "得意先名", "納品先", "得意先"], all_cols)
    c_staff = find_col(["担当者コード", "担当者CD", "担当者名", "担当者", "担当", "営業担当"], all_cols)

    res = pd.DataFrame()
    if c_pcode: res["商品コード"] = df_all[c_pcode].astype(str).str.strip()
    if c_pname:
        res["元商品名"] = df_all[c_pname].astype(str).str.strip()
        res["商品名"] = df_all[c_pname].astype(str).map(clean_product_name)
    if c_cname: res["得意先名"] = df_all[c_cname].astype(str).str.strip().map(clean_cust_name)
    else: res["得意先名"] = ""

    if c_staff: res["担当者"] = df_all[c_staff].map(convert_staff_name).fillna("").replace("nan", "")
    else: res["担当者"] = ""

    res["数量"] = df_all["_auto_qty"].fillna(0).astype(int)
    res["単価"] = df_all["_auto_price"].fillna(0).astype(int)
    res["金額"] = df_all["_auto_amt"].fillna(res["数量"] * res["単価"]).astype(int)
    if "_source_file" in df_all.columns: res["_source_file"] = df_all["_source_file"]

    raw_s = df_all[c_date].astype(str).str.strip() if c_date else pd.Series([""] * len(df_all))
    clean_date = raw_s.str.replace(r'\.0$', '', regex=True).str.replace(r'[-/.\s]', '', regex=True)
    if "_fallback_date" in df_all.columns:
        clean_date = clean_date.where(clean_date.str.len() == 8, df_all["_fallback_date"].astype(str))

    dt = pd.to_datetime(clean_date, format="%Y%m%d", errors="coerce").fillna(pd.to_datetime(raw_s, errors="coerce"))
    res["日付_dt"] = dt
    res["日付"] = dt.dt.strftime('%Y/%m/%d').fillna("")
    res["年月"] = dt.dt.strftime('%Y年%m月').fillna("")
    if "商品名" in res.columns: res = res[~res["商品名"].str.contains("送料", na=False)].copy()
    return res

@st.cache_data(show_spinner=False)
def load_and_prep_data(force_reload=False):
    cost_path, cost_mtime = get_cost_master_file_info()
    cost_master = load_cost_master(cost_path, cost_mtime)

    if not force_reload and os.path.exists(CACHE_FILE):
        try:
            cached_df = pd.read_parquet(CACHE_FILE)
            if not cached_df.empty:
                cached_df["日付_dt"] = pd.to_datetime(cached_df["日付_dt"])
                cached_df["単価原価"] = cached_df["商品コード"].map(lambda c: cost_master.get(str(c).strip(), 0.0))
                cached_df["原価合計"] = (cached_df["数量"] * cached_df["単価原価"]).round().astype(int)
                cached_df["粗利額"] = cached_df["金額"] - cached_df["原価合計"]
                return cached_df, 1, cost_master, 0
        except Exception: pass

    csv_candidates = glob.glob(os.path.join(BASE_SHARE, "**", "*.csv"), recursive=True)
    csv_candidates = sorted(list(set(csv_candidates)))

    yayoi_files = [f for f in csv_candidates if "yayoi" in f.lower()]
    amz_files = [f for f in csv_candidates if "アマゾン売上" in f or "businessreport" in f.lower()]
    if not yayoi_files: yayoi_files = [f for f in csv_candidates if f not in amz_files]

    with ThreadPoolExecutor(max_workers=16) as executor:
        res_list = list(executor.map(_read_single_csv, yayoi_files))
    valid_res = [d for d in res_list if d is not None and not d.empty]
    new_df = _parse_raw_df_to_standard(pd.concat(valid_res, ignore_index=True)) if valid_res else pd.DataFrame()

    amz_res = [_read_amazon_seller_csv(af) for af in amz_files]
    valid_amz = [ad for ad in amz_res if ad is not None and not ad.empty]
    new_amz_df = pd.concat(valid_amz, ignore_index=True) if valid_amz else pd.DataFrame()

    valid_dfs = [d for d in [new_df, new_amz_df] if d is not None and not d.empty]
    df = pd.concat(valid_dfs, ignore_index=True) if valid_dfs else pd.DataFrame()

    DEFAULT_AMZ_SEP_2026 = [
        {"sku": "ABA-501", "qty": 100, "amount": 50000},
        {"sku": "ABA-507", "qty": 50, "amount": 30000}
    ]

    if df.empty or ("アマゾンセラー" not in df.get("得意先名", pd.Series()).values):
        amz_rows = []
        sep_date = pd.Timestamp(2026, 9, 30)
        for r in DEFAULT_AMZ_SEP_2026:
            u_p = int(r["amount"] / r["qty"]) if r["qty"] > 0 else r["amount"]
            amz_rows.append({
                "商品コード": r["sku"], "商品名": r["sku"], "得意先名": "アマゾンセラー", "担当者": "水原",
                "数量": r["qty"], "単価": u_p, "金額": r["amount"], "日付_dt": sep_date,
                "日付": "2026/09/30", "年月": "2026年09月", "_source_file": "default_amz"
            })
        df = pd.concat([df, pd.DataFrame(amz_rows)], ignore_index=True)

    if not df.empty and "担当者" in df.columns and "得意先名" in df.columns:
        valid_staff_df = df[df["担当者"].astype(str).str.strip().ne("")].sort_values("日付_dt")
        if not valid_staff_df.empty:
            cust_current_staff_map = valid_staff_df.groupby("得意先名")["担当者"].last().to_dict()
            df["担当者"] = df["得意先名"].map(lambda c: cust_current_staff_map.get(c, ""))

    if not df.empty and "商品コード" in df.columns:
        df["単価原価"] = df["商品コード"].map(lambda c: cost_master.get(str(c).strip(), 0.0))
        df["原価合計"] = (df["数量"] * df["単価原価"]).round().astype(int)
        df["粗利額"] = df["金額"] - df["原価合計"]

    try:
        if not df.empty: df.to_parquet(CACHE_FILE, index=False)
    except Exception: pass

    return df, len(csv_candidates), cost_master, len(csv_candidates)

def render_universal_table(df_data):
    if df_data is None or df_data.empty:
        st.info("該当するデータがありません。")
        return
    center_cols = {"番号", "順位", "日付", "販売期間", "最終売上日", "初回売上日", "未販売期間"}
    left_cols = {"得意先名", "商品コード", "商品名", "商品番号", "ランク名", "担当者", "客先名", "販売価格帯"}

    header_html = "".join([f"<th style='border:1px solid #cbd5e1; padding:8px 10px; background:#f1f5f9; text-align:{('center' if c in center_cols else ('left' if c in left_cols else 'right'))}; white-space:nowrap;'>{c}</th>" for c in df_data.columns])
    body_rows = []
    for r_i, row in df_data.iterrows():
        bg = "#ffffff" if r_i % 2 == 0 else "#f8fafc"
        cells = []
        for c in df_data.columns:
            val = row[c]
            val_str = "" if pd.isna(val) else str(val)
            align = "center" if c in center_cols else ("left" if c in left_cols else "right")
            if isinstance(val, (int, float)) and c not in center_cols:
                val_str = f"{int(val):+,}" if c in ["増減", "増減額", "目標差額", "粗利増減"] else f"{int(val):,}"

            is_neg = False
            if isinstance(val, (int, float)) and val < 0: is_neg = True
            elif isinstance(val_str, str) and (val_str.startswith("-") or val_str.startswith("▲-")): is_neg = True

            style_extra = "color: #dc2626; font-weight: bold;" if is_neg else "color: #0f172a;"
            cells.append(f"<td style='border:1px solid #e2e8f0; padding:6px 10px; text-align:{align}; font-variant-numeric:tabular-nums; {style_extra} white-space:nowrap;'>{val_str}</td>")
        body_rows.append(f"<tr style='background:{bg};'>{''.join(cells)}</tr>")

    table_html = f"""
    <div style="overflow-x: auto; max-height: 600px; border: 1px solid #cbd5e1; border-radius: 6px; background-color: #ffffff;">
    <table style="width: 100%; border-collapse: collapse; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; font-size: 13px;">
    <thead><tr style="position: sticky; top: 0; z-index: 1;">{header_html}</tr></thead>
    <tbody>{''.join(body_rows)}</tbody>
    </table>
    </div>
    """
    calc_h = min(620, 80 + len(df_data) * 36)
    components.html(table_html, height=calc_h, scrolling=True)

@st.dialog("🏢 得意先別 売れ筋・売上TOP10詳細", width="large")
def show_customer_top10_dialog(cust_name, sub_df, period_str):
    st.markdown(f"#### 🎯 対象得意先: **{cust_name}**")
    st.caption(f"指定期間: {period_str}")
    if sub_df.empty:
        st.info("該当期間の売上データはありません。")
        return

    p_agg = sub_df.groupby(["商品コード", "商品名"]).agg(
        数量=("数量", "sum"), 販売金額=("金額", "sum"), 粗利額=("粗利額", "sum")
    ).reset_index().sort_values(by="販売金額", ascending=False).head(10)

    p_agg["粗利率"] = p_agg.apply(lambda r: f"{(r['粗利額'] / r['販売金額'] * 100):.1f}%" if r['販売金額'] > 0 else "-", axis=1)
    p_agg.insert(0, "順位", range(1, len(p_agg) + 1))
    show_p_df = p_agg[["順位", "商品コード", "商品名", "数量", "販売金額", "粗利額", "粗利率"]]
    
    tot_s = int(p_agg["販売金額"].sum())
    tot_q = int(p_agg["数量"].sum())
    st.markdown(f"<div class='total-card'><b>TOP10商品売上合計:</b> ¥{tot_s:,} &nbsp;|&nbsp; <b>数量合計:</b> {tot_q:,} 個</div>", unsafe_allow_html=True)
    render_universal_table(show_p_df)

# ========================================================
# ★ 営業会議用：上位社 売上比較シート（HTMLインデント除去＆Excel完全調整）
# ========================================================
def _extract_top5_and_other(sub_df):
    if sub_df is None or sub_df.empty:
        return [{"sku": "", "price": "", "qty": "", "amt": ""} for _ in range(5)], 0, 0
    agg = sub_df.groupby("商品コード").agg(
        数量=("数量", "sum"),
        金額=("金額", "sum")
    ).reset_index().sort_values(by="金額", ascending=False)
    total_amt = int(agg["金額"].sum())
    top5_df = agg.head(5)
    top5_sum = int(top5_df["金額"].sum())
    other_amt = max(0, total_amt - top5_sum)
    rows = []
    for _, r in top5_df.iterrows():
        q = int(r["数量"])
        a = int(r["金額"])
        u_p = int(round(a / q)) if q > 0 else 0
        rows.append({"sku": str(r["商品コード"]), "price": u_p, "qty": q, "amt": a})
    while len(rows) < 5:
        rows.append({"sku": "", "price": "", "qty": "", "amt": ""})
    return rows, other_amt, total_amt

def build_meeting_comparison_card_data(cust, cid, cur_m, ly_m, next_m, prev_fy, tgt_y, tgt_m, next_y, next_m_val, f_year, max_dt_val=None):
    c_cur = cur_m[cur_m["得意先名"] == cust] if not cur_m.empty else pd.DataFrame()
    c_ly = ly_m[ly_m["得意先名"] == cust] if not ly_m.empty else pd.DataFrame()
    c_next = next_m[next_m["得意先名"] == cust] if not next_m.empty else pd.DataFrame()
    c_prev_fy = prev_fy[prev_fy["得意先名"] == cust] if not prev_fy.empty else pd.DataFrame()

    ly_rows, ly_other, ly_total = _extract_top5_and_other(c_ly)
    cur_rows, cur_other, cur_total = _extract_top5_and_other(c_cur)
    next_rows, next_other, next_total = _extract_top5_and_other(c_next)
    prev_fy_total = int(c_prev_fy["金額"].sum()) if not c_prev_fy.empty else 0

    yoy_rate = f"{int(round(cur_total / ly_total * 100))}%" if ly_total > 0 else "-"
    next_yoy_rate = f"{int(round(next_total / prev_fy_total * 100))}%" if prev_fy_total > 0 else "0%"

    if not c_next.empty and "日付_dt" in c_next.columns and c_next["日付_dt"].notna().any():
        max_n_dt = c_next["日付_dt"].max()
        as_of_str = f"{next_y}年{next_m_val}月{max_n_dt.day}日現在"
    else:
        as_of_str = f"{next_y}年{next_m_val}月現在"

    return {
        "cust_name": cust,
        "cust_id": cid,
        "ly_label": f"{tgt_y - 1}年{tgt_m}月実績",
        "cur_label": f"{tgt_y}年{tgt_m}月実績",
        "next_label": f"{next_y}年{next_m_val}月売上進捗",
        "prev_fy_label": f"{next_y - 1}年{next_m_val}月度実績",
        "as_of_label": as_of_str,
        "ly_rows": ly_rows, "ly_other": ly_other, "ly_total": ly_total,
        "cur_rows": cur_rows, "cur_other": cur_other, "cur_total": cur_total,
        "yoy_rate": yoy_rate,
        "next_rows": next_rows, "next_other": next_other, "next_total": next_total,
        "prev_fy_total": prev_fy_total, "next_yoy_rate": next_yoy_rate,
        "cur_sales": cur_total, "ly_sales": ly_total, "diff": cur_total - ly_total,
        "next_sales": next_total, "prev_fy_sales": prev_fy_total, "next_diff": next_total - prev_fy_total,
        "tgt_m": tgt_m, "next_m": next_m_val,
        "activity_text": "", "plan_text": ""
    }

def render_meeting_card_html(card):
    def fmt_v(v):
        if v == "" or v is None: return "&nbsp;"
        return f"{int(v):,}" if isinstance(v, (int, float)) else str(v)

    top_rows_list = []
    for i in range(5):
        lr = card["ly_rows"][i]
        cr = card["cur_rows"][i]
        top_rows_list.append(
            f"<tr>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:center; font-size:11px;'>{fmt_v(lr['sku'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(lr['price'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(lr['qty'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(lr['amt'])}</td>"
            f"<td style='border:1px solid #333; border-left:2px solid #111; padding:2px 4px; text-align:center; font-size:11px;'>{fmt_v(cr['sku'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(cr['price'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(cr['qty'])}</td>"
            f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(cr['amt'])}</td>"
            f"</tr>"
        )
    top_rows_html = "".join(top_rows_list)

    act_html = str(card.get("activity_text", "")).replace("\n", "<br>") if card.get("activity_text") else "&nbsp;"
    plan_html = str(card.get("plan_text", "")).replace("\n", "<br>") if card.get("plan_text") else "&nbsp;"

    next_rows_list = []
    for i in range(5):
        nr = card["next_rows"][i]
        if i == 0:
            next_rows_list.append(
                f"<tr>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:center; font-size:11px;'>{fmt_v(nr['sku'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['price'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['qty'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['amt'])}</td>"
                f"<td colspan='4' rowspan='7' style='border:1px solid #333; border-left:2px solid #111; padding:4px 6px; vertical-align:top; font-size:11px; text-align:left;'>{act_html}</td>"
                f"</tr>"
            )
        else:
            next_rows_list.append(
                f"<tr>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:center; font-size:11px;'>{fmt_v(nr['sku'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['price'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['qty'])}</td>"
                f"<td style='border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;'>{fmt_v(nr['amt'])}</td>"
                f"</tr>"
            )
    next_rows_html = "".join(next_rows_list)

    raw_html = f"""
<div style="border:1px solid #94a3b8; padding:10px; background:#ffffff; margin-bottom:14px; font-family:'Yu Gothic','Meiryo',sans-serif; color:#111;">
<div style="display:flex; justify-content:space-between; align-items:baseline; margin-bottom:6px;">
<span style="font-size:15px; font-weight:bold; text-decoration:underline;">{card['cust_name']}</span>
<span style="font-size:12px; color:#333; margin-right:35%;">{card['cust_id']}</span>
</div>
<table style="width:100%; border-collapse:collapse; table-layout:fixed; margin-bottom:2px;">
<thead>
<tr style="background:#f8fafc;">
<th colspan="4" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">{card['ly_label']}</th>
<th colspan="4" style="border:1px solid #333; border-left:2px solid #111; padding:2px 5px; text-align:left; font-size:11px;">{card['cur_label']}</th>
</tr>
<tr style="background:#f1f5f9;">
<th style="border:1px solid #333; padding:2px; font-size:11px; width:14%;">品番</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">単価</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">数量</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:16%;">金額</th>
<th style="border:1px solid #333; border-left:2px solid #111; padding:2px; font-size:11px; width:14%;">品番</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">単価</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">数量</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:16%;">金額</th>
</tr>
</thead>
<tbody>
{top_rows_html}
<tr>
<td colspan="3" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">その他</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['ly_other']:,}</td>
<td colspan="3" style="border:1px solid #333; border-left:2px solid #111; padding:2px 5px; text-align:left; font-size:11px;">その他</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['cur_other']:,}</td>
</tr>
<tr style="font-weight:bold; background:#f8fafc;">
<td colspan="3" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">合計</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['ly_total']:,}</td>
<td colspan="3" style="border:1px solid #333; border-left:2px solid #111; padding:2px 5px; text-align:left; font-size:11px;">合計</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['cur_total']:,}</td>
</tr>
</tbody>
</table>
<div style="text-align:right; font-size:11px; font-weight:bold; margin-bottom:8px; padding-right:4px;">
昨対比 &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; {card['yoy_rate']}
</div>
<table style="width:100%; border-collapse:collapse; table-layout:fixed;">
<thead>
<tr style="background:#f8fafc;">
<th colspan="4" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">{card['next_label']}</th>
<th colspan="4" style="border:1px solid #333; border-left:2px solid #111; padding:2px 5px; text-align:left; font-size:11px;">この1ヶ月の主な営業活動</th>
</tr>
<tr style="background:#f1f5f9;">
<th style="border:1px solid #333; padding:2px; font-size:11px; width:14%;">品番</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">単価</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:10%;">数量</th>
<th style="border:1px solid #333; padding:2px; font-size:11px; width:16%;">金額</th>
<th colspan="4" style="border:1px solid #333; border-left:2px solid #111; padding:0; height:2px;"></th>
</tr>
</thead>
<tbody>
{next_rows_html}
<tr>
<td colspan="3" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">その他</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['next_other']:,}</td>
</tr>
<tr style="font-weight:bold; background:#f8fafc;">
<td colspan="3" style="border:1px solid #333; padding:2px 5px; text-align:left; font-size:11px;">合計</td>
<td style="border:1px solid #333; padding:2px 4px; text-align:right; font-size:11px;">{card['next_total']:,}</td>
</tr>
<tr>
<td colspan="4" style="border:1px solid #333; padding:4px 6px; vertical-align:top; height:68px; font-size:11px;">
<div style="color:#64748b; font-size:10px;">【今後の予定】</div>
<div>{plan_html}</div>
</td>
<td colspan="4" style="border:1px solid #333; border-left:2px solid #111; padding:4px 6px; vertical-align:top; font-size:11px;">
<div style="display:flex; justify-content:space-between;"><span>{card['prev_fy_label']}</span><span>{card['prev_fy_total']:,}</span></div>
<div style="display:flex; justify-content:space-between;"><span>{card['as_of_label']}</span><span>{card['next_total']:,}</span></div>
<div style="display:flex; justify-content:space-between; font-weight:bold;"><span>昨対進捗</span><span>{card['next_yoy_rate']}</span></div>
<div style="font-size:10px; color:#475569; margin-top:2px;">残り営業日数13日（3/16日程消化19%）</div>
</td>
</tr>
</tbody>
</table>
</div>
"""
    return "\n".join(line.strip() for line in raw_html.splitlines() if line.strip())

def generate_meeting_comparison_excel(cards_data, tgt_ym):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "上位社売上比較シート"
    font_b = Font(name="Yu Gothic", size=10, bold=True)
    font_m = Font(name="Yu Gothic", size=9)
    thin = Side(style='thin', color='333333')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_h = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    col_widths = {
        1: 15, 2: 9, 3: 9, 4: 14, 5: 15, 6: 9, 7: 9, 8: 14,
        9: 4,
        10: 15, 11: 9, 12: 9, 13: 14, 14: 15, 15: 9, 16: 9, 17: 14
    }
    for col_idx, width_val in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width_val

    for idx, c in enumerate(cards_data):
        grid_r = idx // 2
        grid_c = idx % 2
        base_r = 2 + grid_r * 26
        base_c = 1 + grid_c * 9

        ws.cell(row=base_r, column=base_c, value=c["cust_name"]).font = Font(name="Yu Gothic", size=11, bold=True, underline="single")
        ws.cell(row=base_r, column=base_c + 2, value=c["cust_id"]).font = font_m

        r_top = base_r + 1
        ws.cell(row=r_top, column=base_c, value=c["ly_label"]).font = font_b
        ws.cell(row=r_top, column=base_c + 4, value=c["cur_label"]).font = font_b

        headers = ["品番", "単価", "数量", "金額", "品番", "単価", "数量", "金額"]
        for h_i, h_t in enumerate(headers):
            cell = ws.cell(row=r_top + 1, column=base_c + h_i, value=h_t)
            cell.font = font_b; cell.border = bd; cell.fill = fill_h
            cell.alignment = Alignment(horizontal="center")

        for i in range(5):
            curr_row = r_top + 2 + i
            lr = c["ly_rows"][i]; cr = c["cur_rows"][i]
            vals = [lr["sku"], lr["price"], lr["qty"], lr["amt"], cr["sku"], cr["price"], cr["qty"], cr["amt"]]
            for v_i, v_val in enumerate(vals):
                cell = ws.cell(row=curr_row, column=base_c + v_i, value=v_val if v_val != "" else None)
                cell.font = font_m; cell.border = bd
                if v_i in [1, 2, 3, 5, 6, 7] and isinstance(v_val, (int, float)):
                    cell.number_format = '#,##0'

        r_oth = r_top + 7
        ws.cell(row=r_oth, column=base_c, value="その他").font = font_m
        ws.cell(row=r_oth, column=base_c + 3, value=c["ly_other"]).number_format = '#,##0'
        ws.cell(row=r_oth, column=base_c + 3).font = font_m
        ws.cell(row=r_oth, column=base_c + 4, value="その他").font = font_m
        ws.cell(row=r_oth, column=base_c + 7, value=c["cur_other"]).number_format = '#,##0'
        ws.cell(row=r_oth, column=base_c + 7).font = font_m
        for h_i in range(8): ws.cell(row=r_oth, column=base_c + h_i).border = bd

        r_tot = r_top + 8
        ws.cell(row=r_tot, column=base_c, value="合計").font = font_b
        ws.cell(row=r_tot, column=base_c + 3, value=c["ly_total"]).number_format = '#,##0'
        ws.cell(row=r_tot, column=base_c + 3).font = font_b
        ws.cell(row=r_tot, column=base_c + 4, value="合計").font = font_b
        ws.cell(row=r_tot, column=base_c + 7, value=c["cur_total"]).number_format = '#,##0'
        ws.cell(row=r_tot, column=base_c + 7).font = font_b
        for h_i in range(8): ws.cell(row=r_tot, column=base_c + h_i).border = bd

        ws.cell(row=r_tot + 1, column=base_c + 5, value="昨対比").font = font_b
        c_yoy = ws.cell(row=r_tot + 1, column=base_c + 7, value=c["yoy_rate"])
        c_yoy.font = font_b; c_yoy.alignment = Alignment(horizontal="right")

        r_nx = r_tot + 2
        ws.cell(row=r_nx, column=base_c, value=c["next_label"]).font = font_b
        ws.cell(row=r_nx, column=base_c + 4, value="この1ヶ月の主な営業活動").font = font_b
        for h_i, h_t in enumerate(["品番", "単価", "数量", "金額"]):
            cell = ws.cell(row=r_nx + 1, column=base_c + h_i, value=h_t)
            cell.font = font_b; cell.border = bd; cell.fill = fill_h
            cell.alignment = Alignment(horizontal="center")

        for i in range(5):
            curr_row = r_nx + 2 + i
            nr = c["next_rows"][i]
            for v_i, v_val in enumerate([nr["sku"], nr["price"], nr["qty"], nr["amt"]]):
                cell = ws.cell(row=curr_row, column=base_c + v_i, value=v_val if v_val != "" else None)
                cell.font = font_m; cell.border = bd
                if v_i in [1, 2, 3] and isinstance(v_val, (int, float)):
                    cell.number_format = '#,##0'

        ws.merge_cells(start_row=r_nx + 1, start_column=base_c + 4, end_row=r_nx + 8, end_column=base_c + 7)
        act_c = ws.cell(row=r_nx + 1, column=base_c + 4, value=c.get("activity_text", ""))
        act_c.font = font_m; act_c.alignment = Alignment(vertical="top", wrap_text=True)
        for rr in range(r_nx + 1, r_nx + 9):
            for cc in range(base_c + 4, base_c + 8):
                ws.cell(row=rr, column=cc).border = bd

        r_nx_oth = r_nx + 7
        ws.cell(row=r_nx_oth, column=base_c, value="その他").font = font_m
        ws.cell(row=r_nx_oth, column=base_c + 3, value=c["next_other"]).number_format = '#,##0'
        for h_i in range(4): ws.cell(row=r_nx_oth, column=base_c + h_i).border = bd

        r_nx_tot = r_nx + 8
        ws.cell(row=r_nx_tot, column=base_c, value="合計").font = font_b
        ws.cell(row=r_nx_tot, column=base_c + 3, value=c["next_total"]).number_format = '#,##0'
        for h_i in range(4): ws.cell(row=r_nx_tot, column=base_c + h_i).border = bd

        r_bot = r_nx + 9
        ws.merge_cells(start_row=r_bot, start_column=base_c, end_row=r_bot + 3, end_column=base_c + 3)
        pln_c = ws.cell(row=r_bot, column=base_c, value=f"【今後の予定】\n{c.get('plan_text', '')}")
        pln_c.font = font_m; pln_c.alignment = Alignment(vertical="top", wrap_text=True)
        for rr in range(r_bot, r_bot + 4):
            for cc in range(base_c, base_c + 4):
                ws.cell(row=rr, column=cc).border = bd

        ws.cell(row=r_bot, column=base_c + 4, value=c["prev_fy_label"]).font = font_m
        c_pfy = ws.cell(row=r_bot, column=base_c + 7, value=c["prev_fy_total"])
        c_pfy.number_format = '#,##0'; c_pfy.font = font_m

        ws.cell(row=r_bot + 1, column=base_c + 4, value=c["as_of_label"]).font = font_m
        c_nxt = ws.cell(row=r_bot + 1, column=base_c + 7, value=c["next_total"])
        c_nxt.number_format = '#,##0'; c_nxt.font = font_m

        ws.cell(row=r_bot + 2, column=base_c + 4, value="昨対進捗").font = font_b
        c_nyoy = ws.cell(row=r_bot + 2, column=base_c + 7, value=c["next_yoy_rate"])
        c_nyoy.font = font_b; c_nyoy.alignment = Alignment(horizontal="right")

        ws.cell(row=r_bot + 3, column=base_c + 4, value="残り営業日数13日（3/16日程消化19%）").font = Font(name="Yu Gothic", size=8)
        for rr in range(r_bot, r_bot + 4):
            for cc in range(base_c + 4, base_c + 8):
                ws.cell(row=rr, column=cc).border = bd

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def generate_meeting_top_cust_excel(res_df, title_str):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        res_df.to_excel(writer, index=False, sheet_name="上位客先実績")
    return out.getvalue()

# ========================================================
# ★ データロード＆基本日付変数初期化
# ========================================================
df, file_count, cost_master_dict, diff_cnt = load_and_prep_data(force_reload=False)

member_p_path, member_p_mtime = get_member_price_file_info()
member_cust_sku_map, member_default_min_map = load_member_price_master_v2(member_p_path, member_p_mtime)

jan_path, jan_mtime = get_jan_file_info()
jan_master_dict, jan_title_dict, jan_std_price_dict = load_jan_master_v2(jan_path, jan_mtime)

cust_path, cust_mtime = get_customer_file_info()
customer_master_df = load_customer_master(cust_path, cust_mtime)

hanko_bytes, hanko_file_path = get_hanko_bytes()

col_status, col_refresh = st.columns([4, 1.2])
with col_refresh:
    if st.button("🔄 データを最新に再読み込み"):
        st.cache_data.clear()
        if os.path.exists(CACHE_FILE): os.remove(CACHE_FILE)
        with st.spinner("最新データを高速読み込み中..."):
            df, file_count, cost_master_dict, diff_cnt = load_and_prep_data(force_reload=True)
            member_cust_sku_map, member_default_min_map = load_member_price_master_v2(member_p_path, member_p_mtime)
            jan_master_dict, jan_title_dict, jan_std_price_dict = load_jan_master_v2(jan_path, jan_mtime)
            customer_master_df = load_customer_master(cust_path, cust_mtime)
            hanko_bytes, hanko_file_path = get_hanko_bytes()
        st.success("最新データに更新しました！")
        time.sleep(0.5)
        st.rerun()

if df.empty:
    st.error(f"売上CSVデータが見つかりませんでした。`{BASE_SHARE}` フォルダを確認してください。")
    st.stop()

valid_ym = sorted([x for x in df["年月"].dropna().unique() if x and str(x).strip() != ""])
latest_ym = valid_ym[-1] if valid_ym else "2026年01月"
desc_valid_ym = sorted(valid_ym, reverse=True)

valid_dts = df["日付_dt"].dropna()
min_date = valid_dts.min().date() if not valid_dts.empty else datetime(2022, 1, 1).date()
max_date = valid_dts.max().date() if not valid_dts.empty else datetime.now().date()

curr_y = int(latest_ym[:4]) if len(latest_ym) >= 4 else 2026
curr_m = int(latest_ym[5:7]) if len(latest_ym) >= 7 else 1
fiscal_year = curr_y if curr_m >= 7 else curr_y - 1

col_top1, col_top2 = st.columns([1.6, 2])
with col_top1:
    view_type = st.selectbox("表示モード", [
        '🧮 簡易見積＆利益シミュレーター',
        '🏢 得意先別ランキング',
        '📦 商品別売上ランキング',
        '📊 営業会議用：上位社 売上比較シート',
        '📑 営業会議用：上位客先別売上実績',
        '🏷️ ランク価格マスタ一覧',
        '👤 担当者別 売上＆目標達成分析',
        '💤 滞留・未販売商品分析（売れてない商品）',
        '✨ 新商品売上分析（直近1年）',
        '📈 商品別 前年実績比較',
        '📊 得意先別 前年実績比較',
        '🧩 得意先×商品 マトリクス',
        '📋 明細一覧'
    ])
with col_top2:
    if view_type == '🧮 簡易見積＆利益シミュレーター':
        st.write("")
        jan_status = f"JAN: **{len(jan_master_dict):,} 件**" if jan_master_dict else "⚠️ JAN未検出"
        hanko_status = "角印: **読込完了**" if hanko_bytes else "⚠️ 印鑑未検出"
        member_status = f"加盟店価格: **{len(member_default_min_map):,} 件**" if member_default_min_map else "⚠️ 加盟店価格未検出"
        st.info(f"📁 原価: **{len(cost_master_dict):,} 件** | {member_status} | {jan_status} | {hanko_status}")
    else:
        quick_choice = st.radio("期間ショートカット", ['当月', '前月', '当期累計', '前期累計', '全データ累計', '手動指定'], horizontal=True, key="period_shortcut_radio_top")

# ========================================================
# ★ 🧮 簡易見積＆利益シミュレーター（完全維持）
# ========================================================
if view_type == '🧮 簡易見積＆利益シミュレーター':
    code_to_name_map = df.dropna(subset=["商品コード", "商品名"]).drop_duplicates(subset=["商品コード"]).set_index("商品コード")["商品名"].to_dict()
    for k_t, v_t in jan_title_dict.items():
        sk_trim = str(k_t).strip()
        if sk_trim not in code_to_name_map or not code_to_name_map[sk_trim]:
            code_to_name_map[sk_trim] = str(v_t).strip()

    raw_prod_title_map = {}
    if "元商品名" in df.columns:
        raw_prod_title_map = df.dropna(subset=["商品コード", "元商品名"]).drop_duplicates(subset=["商品コード"]).set_index("商品コード")["元商品名"].to_dict()
    for k_t, v_t in jan_title_dict.items():
        sk_trim = str(k_t).strip()
        if sk_trim not in raw_prod_title_map:
            raw_prod_title_map[sk_trim] = str(v_t).strip()

    raw_calc_codes = sorted(list(set(
        [str(c).strip() for c in code_to_name_map.keys()] +
        [str(c).strip() for c in cost_master_dict.keys()] +
        [str(c).strip() for c in jan_master_dict.keys()] +
        [str(c).strip() for c in jan_title_dict.keys()] +
        [str(c).strip() for c in member_default_min_map.keys()]
    )))
    if "ABA-40" in raw_calc_codes: raw_calc_codes.remove("ABA-40")
    if "" in raw_calc_codes: raw_calc_codes.remove("")

    parent_sku_candidates = set()
    for c in raw_calc_codes:
        for other in raw_calc_codes:
            if other != c and other.startswith(c) and len(other) > len(c):
                parent_sku_candidates.add(c)
                break
    all_calc_codes = [c for c in raw_calc_codes if c not in parent_sku_candidates]

    STANDARD_RANKS = [
        "ロット単価",
        "1.05倍", "1.07倍", "1.1倍", "1.15倍", "1.2倍", "1.25倍",
        "1.3倍", "1.35倍", "1.4倍", "1.45倍", "1.5倍", "0.9倍"
    ]

    RANK_MULTIPLIERS = {
        "1.05倍": 1.05, "1.07倍": 1.07, "1.1倍": 1.10, "1.15倍": 1.15,
        "1.2倍": 1.20, "1.25倍": 1.25, "1.3倍": 1.30, "1.35倍": 1.35,
        "1.4倍": 1.40, "1.45倍": 1.45, "1.5倍": 1.50, "0.9倍": 0.90
    }

    def get_sku_details(sku_code, rank_selection="ロット単価", cust_comp_name=""):
        s_trim = str(sku_code).strip()
        matched_key = s_trim
        if s_trim not in jan_title_dict and s_trim not in cost_master_dict and s_trim not in member_default_min_map:
            candidates = [k for k in jan_title_dict.keys() if str(k).strip().startswith(s_trim) or s_trim in str(k)]
            if candidates:
                matched_key = str(candidates[0]).strip()

        p_name = jan_title_dict.get(matched_key, "")
        if not p_name or p_name == "nan": 
            p_name = code_to_name_map.get(matched_key, s_trim)
        
        p_jan = jan_master_dict.get(matched_key, "")
        p_cost = int(cost_master_dict.get(matched_key, 0.0))
        
        norm_cust = re.sub(r'[\s ]+', '', str(cust_comp_name)) if cust_comp_name else ""
        min_p = member_default_min_map.get(matched_key, 0)

        comp_lot_p = 0
        if cust_comp_name and (cust_comp_name, matched_key) in member_cust_sku_map:
            comp_lot_p = member_cust_sku_map[(cust_comp_name, matched_key)]
        elif norm_cust and (norm_cust, matched_key) in member_cust_sku_map:
            comp_lot_p = member_cust_sku_map[(norm_cust, matched_key)]

        clean_r = str(rank_selection).strip() if rank_selection else "ロット単価"
        p_price = 0
        actual_rank_applied = clean_r

        if clean_r in RANK_MULTIPLIERS:
            if min_p > 0:
                p_price = int(round(min_p * RANK_MULTIPLIERS[clean_r]))
            else:
                p_price = 99999
        elif clean_r == "ロット単価" or clean_r == "":
            if comp_lot_p > 0:
                p_price = comp_lot_p
            elif min_p > 0:
                p_price = min_p
            else:
                p_price = 99999

        if p_price == 0: p_price = 99999
        return p_name, p_jan, p_price, p_cost, min_p, actual_rank_applied

    if "saved_estimates" not in st.session_state:
        st.session_state.saved_estimates = load_saved_quotes()

    est_staff_list = list(STAFF_MAP.values()) + ["（すべて表示）"]
    col_menu_left, col_menu_right = st.columns([1.6, 2.0])
    with col_menu_right:
        est_staff = st.selectbox(
            "👤 担当者（誰が作業するか選択：バッティング防止）",
            est_staff_list,
            index=7 if "水原" in est_staff_list else 0,
            key="est_staff_selector_box_v3"
        )

    s_key = est_staff
    k_items = f"quotation_items_{s_key}"
    k_ver = f"editor_version_{s_key}"
    k_comp = f"in_cust_comp_{s_key}"
    k_person = f"in_cust_person_{s_key}"
    k_post = f"in_cust_post_{s_key}"
    k_addr = f"in_cust_addr_{s_key}"
    k_s_comp = f"in_ship_comp_{s_key}"
    k_s_person = f"in_ship_person_{s_key}"
    k_s_post = f"in_ship_post_{s_key}"
    k_s_addr = f"in_ship_addr_{s_key}"
    k_s_tel = f"in_ship_tel_{s_key}"
    k_sel_c = f"sel_cust_choice_{s_key}"

    if k_items not in st.session_state: st.session_state[k_items] = []
    if k_ver not in st.session_state: st.session_state[k_ver] = 0
    for k_f in [k_comp, k_person, k_post, k_addr, k_s_comp, k_s_person, k_s_post, k_s_addr, k_s_tel]:
        if k_f not in st.session_state: st.session_state[k_f] = ""

    if "load_target_estimate" in st.session_state and st.session_state["load_target_estimate"]:
        loaded_est = st.session_state.pop("load_target_estimate")
        st.session_state[k_comp] = loaded_est.get("cust_comp", "")
        st.session_state[k_person] = loaded_est.get("cust_person", "")
        st.session_state[k_post] = loaded_est.get("cust_post", "")
        st.session_state[k_addr] = loaded_est.get("cust_addr", "")
        st.session_state[k_s_comp] = loaded_est.get("ship_comp", "")
        st.session_state[k_s_person] = loaded_est.get("ship_person", "")
        st.session_state[k_s_post] = loaded_est.get("ship_post", "")
        st.session_state[k_s_addr] = loaded_est.get("ship_addr", "")
        st.session_state[k_s_tel] = loaded_est.get("ship_tel", "")
        st.session_state[k_sel_c] = loaded_est.get("cust_comp", "（手動入力）")
        st.session_state[k_items] = loaded_est.get("items", [])
        st.session_state[k_ver] += 1

    staff_Filter_list = [
        e for e in st.session_state.saved_estimates
        if est_staff == "（すべて表示）" or e.get("staff", "") == est_staff or (est_staff in str(e.get("hitokoto", "")))
    ]

    with col_menu_left:
        tab_choice = st.radio(
            "見積機能メニュー",
            ["📝 新規見積作成・シミュレーター", f"📋 保存済み見積一覧・スマレジ注文CSV出力（{len(staff_Filter_list)}件）"],
            horizontal=True,
            key="est_active_tab_v3"
        )

    if tab_choice.startswith("📝"):
        st.markdown(f"##### 🏢 提出先（得意先）の選択 &nbsp;<span style='font-size:13px; color:#2563eb;'>［作業担当: {est_staff}］</span>", unsafe_allow_html=True)
        cust_options = ["（手動入力）"]
        cust_lookup = {}
        if not customer_master_df.empty:
            c_name_col = "加盟店名" if "加盟店名" in customer_master_df.columns else customer_master_df.columns[0]
            for _, c_row in customer_master_df.iterrows():
                c_name = str(c_row[c_name_col]).strip()
                if c_name and c_name != "nan":
                    cust_options.append(c_name)
                    cust_lookup[c_name] = c_row

        def on_customer_change():
            selected = st.session_state.get(k_sel_c, "（手動入力）")
            if selected == "（手動入力）" or selected not in cust_lookup:
                for k_f in [k_comp, k_person, k_post, k_addr, k_s_comp, k_s_person, k_s_post, k_s_addr, k_s_tel]:
                    st.session_state[k_f] = ""
            else:
                row_c = cust_lookup[selected]
                sei = str(row_c.get("責任者姓", "")).strip() if pd.notna(row_c.get("責任者姓", "")) else ""
                mei = str(row_c.get("責任者名", "")).strip() if pd.notna(row_c.get("責任者名", "")) else ""
                p_name = f"{sei} {mei}".strip()
                if p_name and "様" not in p_name: p_name += " 様"

                p_code = str(row_c.get("郵便番号", "")).replace(".0", "").strip() if pd.notna(row_c.get("郵便番号", "")) else ""
                pref = str(row_c.get("都道府県", "")).strip() if pd.notna(row_c.get("都道府県", "")) else ""
                a1 = str(row_c.get("住所１", "")).strip() if pd.notna(row_c.get("住所１", "")) else ""
                a2 = str(row_c.get("住所２", "")).strip() if pd.notna(row_c.get("住所２", "")) else ""
                full_addr = f"{pref}{a1} {a2}".strip()
                tel = str(row_c.get("電話番号", "")).strip() if pd.notna(row_c.get("電話番号", "")) else ""

                st.session_state[k_comp] = selected
                st.session_state[k_person] = p_name
                st.session_state[k_post] = p_code
                st.session_state[k_addr] = full_addr
                st.session_state[k_s_comp] = selected
                st.session_state[k_s_person] = p_name
                st.session_state[k_s_post] = p_code
                st.session_state[k_s_addr] = full_addr
                st.session_state[k_s_tel] = tel

                if st.session_state[k_items]:
                    for itm in st.session_state[k_items]:
                        if itm.get("適用ランク") == "ロット単価":
                            _, _, new_p, _, new_min, _ = get_sku_details(itm["商品コード"], "ロット単価", selected)
                            itm["見積単価"] = new_p
                            itm["最低下代"] = new_min
                    st.session_state[k_ver] += 1

        st.selectbox("🔍 加盟店・取引先を選択（検索可能）", cust_options, key=k_sel_c, on_change=on_customer_change)

        col_addr1, col_addr2 = st.columns(2)
        with col_addr1:
            st.markdown("###### 📄 宛先情報（見積書 表面）")
            q_cust_comp = st.text_input("宛先 会社名", key=k_comp)
            q_cust_person = st.text_input("宛先 担当者名", key=k_person)
            q_cust_post = st.text_input("宛先 郵便番号", key=k_post)
            q_cust_addr = st.text_input("宛先 住所", key=k_addr)

        with col_addr2:
            st.markdown("###### 🚚 お届け先情報（送り先）")
            ship_cust_comp = st.text_input("お届け先 会社名", key=k_s_comp)
            ship_cust_person = st.text_input("お届け先 担当者名", key=k_s_person)
            ship_cust_post = st.text_input("お届け先 郵便番号", key=k_s_post)
            ship_cust_addr = st.text_input("お届け先 住所", key=k_s_addr)
            ship_cust_tel = st.text_input("お届け先 TEL", key=k_s_tel)

        hq1, hq2, hq3, hq4 = st.columns(4)
        with hq1: q_no = st.text_input("見積番号", value=f"{datetime.now().strftime('%Y%m%d')}-01", key=f"q_no_{s_key}")
        with hq2: q_date = st.date_input("お見積日", value=datetime.now().date(), key=f"q_date_{s_key}")
        with hq3: q_exp_date = st.date_input("見積有効期限", value=(datetime.now() + timedelta(days=30)).date(), key=f"q_exp_{s_key}")
        with hq4: q_payment = st.selectbox("お支払い方法", ["銀行振込", "代金引換", "現行通り", "掛け払い（月末締め翌月末払い）"], index=0, key=f"q_pay_{s_key}")

        hq5, hq6, hq7 = st.columns([2, 2, 1])
        with hq5: q_memo = st.text_input("伝達事項 / 発注NO等（CSV要望欄に連動）", placeholder="例: 発注NO.2352640", key=f"q_memo_{s_key}")
        with hq6: q_staff_memo = st.text_input("営業担当・社内備考", value=f"営業担当：{est_staff}", key=f"q_smemo_{s_key}")
        with hq7: q_shipping_fee = st.number_input("送料(税抜)", min_value=0, value=0, step=100, key=f"q_ship_{s_key}")

        st.markdown("---")
        st.markdown("##### 📑 複数品目のまとめ見積リスト（見積書作成）")

        def apply_rank_to_all_items(target_rank, current_cust=""):
            if st.session_state[k_items]:
                for item in st.session_state[k_items]:
                    _, _, new_p, _, _, applied_lbl = get_sku_details(item["商品コード"], target_rank, current_cust)
                    item["適用ランク"] = applied_lbl
                    item["見積単価"] = int(new_p)
                    item["合計金額"] = int(item["数量"]) * int(new_p)
                    item["粗利額"] = item["合計金額"] - (int(item["数量"]) * int(item["単価原価"]))
                    item["粗利率"] = f"{(item['粗利額'] / item['合計金額'] * 100):.1f}%" if item["合計金額"] > 0 else "-"
                st.session_state[k_ver] += 1

        c_btn_del, c_btn_clear, c_rank_box, _ = st.columns([1.5, 1.2, 2.0, 1.3])
        trigger_delete_selected = False
        with c_btn_del:
            if st.button("🗑️ 選択した行を削除", key=f"btn_del_{s_key}"): trigger_delete_selected = True
        with c_btn_clear:
            if st.button("全クリア", key=f"btn_clr_{s_key}"):
                st.session_state[k_items] = []
                st.session_state[k_ver] += 1
                st.rerun()
        with c_rank_box:
            rk_key = f"in_cust_rank_{s_key}"
            st.selectbox("🏷️ 価格ランク一括指定", STANDARD_RANKS, key=rk_key, on_change=lambda: apply_rank_to_all_items(st.session_state[rk_key], q_cust_comp), label_visibility="collapsed")

        cols_order = ["選択", "商品コード", "商品名", "JANコード", "最低下代", "適用ランク", "数量", "見積単価", "合計金額", "単価原価", "粗利額", "粗利率"]

        if st.session_state[k_items]:
            raw_items_df = pd.DataFrame(st.session_state[k_items])
            for c in cols_order:
                if c not in raw_items_df.columns:
                    if c == "選択": raw_items_df[c] = False
                    elif c == "適用ランク": raw_items_df[c] = "ロット単価"
                    elif c == "粗利率": raw_items_df[c] = "-"
                    else: raw_items_df[c] = 0
            raw_items_df["数量"] = raw_items_df["数量"].apply(lambda x: max(1, int(clean_num(x))))
            raw_items_df["最低下代"] = raw_items_df["最低下代"].apply(lambda x: int(clean_num(x)))
            raw_items_df["見積単価"] = raw_items_df["見積単価"].apply(lambda x: int(clean_num(x)))
            raw_items_df["単価原価"] = raw_items_df["単価原価"].apply(lambda x: int(clean_num(x)))
            raw_items_df["合計金額"] = raw_items_df["数量"] * raw_items_df["見積単価"]
            raw_items_df["粗利額"] = raw_items_df["合計金額"] - (raw_items_df["数量"] * raw_items_df["単価原価"])
            raw_items_df["粗利率"] = raw_items_df.apply(lambda r: f"{(r['粗利額'] / r['合計金額'] * 100):.1f}%" if r['合計金額'] > 0 else "-", axis=1)
            raw_items_df = raw_items_df[cols_order]
        else:
            raw_items_df = pd.DataFrame([{
                "選択": False, "商品コード": "", "商品名": "", "JANコード": "", "最低下代": 0,
                "適用ランク": "ロット単価", "数量": 1, "見積単価": 0, "合計金額": 0, "単価原価": 0, "粗利額": 0, "粗利率": "-"
            }], columns=cols_order)

        edited_df = st.data_editor(
            raw_items_df,
            column_config={
                "選択": st.column_config.CheckboxColumn("選択", default=False, width="small"),
                "商品コード": st.column_config.SelectboxColumn("商品コード（検索・選択）", options=all_calc_codes, required=True, width="medium"),
                "商品名": st.column_config.TextColumn("商品名", width="large"),
                "JANコード": st.column_config.TextColumn("JANコード"),
                "最低下代": st.column_config.NumberColumn("最低下代", disabled=True, format="¥%d"),
                "適用ランク": st.column_config.SelectboxColumn("適用ランク (▼)", options=STANDARD_RANKS, width="medium"),
                "数量": st.column_config.NumberColumn("数量", min_value=1, step=1),
                "見積単価": st.column_config.NumberColumn("見積単価（円）", min_value=0, step=10, format="¥%d"),
                "合計金額": st.column_config.NumberColumn("合計金額", disabled=True, format="¥%d"),
                "単価原価": st.column_config.NumberColumn("単価原価", format="¥%d"),
                "粗利額": st.column_config.NumberColumn("粗利額", disabled=True, format="¥%d"),
                "粗利率": st.column_config.TextColumn("粗利率", disabled=True)
            },
            num_rows="dynamic",
            use_container_width=True,
            hide_index=True,
            key=f"quote_editor_{s_key}_{st.session_state[k_ver]}"
        )

        updated_records = []
        needs_rerun = False

        for idx_e, row_e in edited_df.iterrows():
            is_selected = bool(row_e.get("選択", False))
            if trigger_delete_selected and is_selected:
                needs_rerun = True
                continue

            c_code = str(row_e["商品コード"]).strip() if pd.notna(row_e["商品コード"]) else ""
            if not c_code or c_code == "None" or c_code == "ABA-40": continue

            prev_rec = st.session_state[k_items][idx_e] if idx_e < len(st.session_state[k_items]) else {}
            prev_code = str(prev_rec.get("商品コード", "")).strip()
            prev_rank = str(prev_rec.get("適用ランク", "ロット単価")).strip()
            prev_qty = int(clean_num(prev_rec.get("数量", 1)))
            prev_price = int(clean_num(prev_rec.get("見積単価", 0)))

            cur_qty = max(1, int(clean_num(row_e.get("数量", 1))))
            cur_rank = str(row_e.get("適用ランク", "ロット単価")).strip() if pd.notna(row_e.get("適用ランク")) else "ロット単価"
            cur_price = int(clean_num(row_e.get("見積単価", 0)))
            cur_cost = int(clean_num(row_e.get("単価原価", 0)))
            cur_name = str(row_e.get("商品名", "")).strip() if pd.notna(row_e.get("商品名")) else ""
            cur_jan = str(row_e.get("JANコード", "")).strip() if pd.notna(row_e.get("JANコード")) else ""
            cur_min = int(clean_num(row_e.get("最低下代", 0)))

            if c_code != prev_code or not cur_name or cur_name == "None" or cur_price == 0:
                auto_name, auto_jan, auto_price, auto_cost, auto_min, applied_lbl = get_sku_details(c_code, cur_rank if cur_rank != "None" else "ロット単価", q_cust_comp)
                cur_name = auto_name
                cur_jan = str(auto_jan)
                cur_price = int(auto_price)
                cur_cost = int(auto_cost)
                cur_min = int(auto_min)
                cur_rank = applied_lbl
                needs_rerun = True
            elif cur_rank != prev_rank:
                _, _, rank_p, _, _, applied_lbl = get_sku_details(c_code, cur_rank, q_cust_comp)
                cur_price = int(rank_p)
                cur_rank = applied_lbl
                needs_rerun = True
            elif cur_qty != prev_qty or cur_price != prev_price:
                needs_rerun = True

            new_total = cur_qty * cur_price
            new_profit = new_total - (cur_qty * cur_cost)
            new_margin_str = f"{(new_profit / new_total * 100):.1f}%" if new_total > 0 else "-"

            updated_records.append({
                "選択": False if trigger_delete_selected else is_selected,
                "商品コード": c_code,
                "商品名": cur_name,
                "JANコード": cur_jan,
                "最低下代": cur_min,
                "適用ランク": cur_rank,
                "数量": cur_qty,
                "見積単価": cur_price,
                "合計金額": new_total,
                "単価原価": cur_cost,
                "粗利額": new_profit,
                "粗利率": new_margin_str
            })

        if trigger_delete_selected or needs_rerun or len(updated_records) != len(st.session_state[k_items]):
            st.session_state[k_items] = updated_records
            st.session_state[k_ver] += 1
            st.rerun()
        else:
            st.session_state[k_items] = updated_records

        if updated_records:
            calc_df = pd.DataFrame(updated_records)
            q_tot_qty = int(calc_df["数量"].sum())
            q_tot_unit_price = int(calc_df["見積単価"].sum())
            q_tot_sales = int(calc_df["合計金額"].sum())
            q_tot_profit = int(calc_df["粗利額"].sum())
            q_tot_margin = (q_tot_profit / q_tot_sales * 100) if q_tot_sales > 0 else 0.0
            q_tax = int(round((q_tot_sales + q_shipping_fee) * 0.10))
            q_grand_total = q_tot_sales + q_shipping_fee + q_tax

            st.markdown(f"""
            <div class="total-card" style="margin-top:10px;">
                <b>数量合計:</b> <span style="font-size:16px; font-weight:bold; color:#1a73e8;">{q_tot_qty:,} 個</span> &nbsp;|&nbsp;
                <b>税抜商品計:</b> <span style="font-size:16px; font-weight:bold;">¥{q_tot_sales:,}</span> &nbsp;|&nbsp;
                <b>送料(税抜):</b> ¥{q_shipping_fee:,} &nbsp;|&nbsp;
                <b>消費税:</b> ¥{q_tax:,} &nbsp;|&nbsp;
                <b>税込合計額:</b> <span style="color:#d93025; font-size:18px; font-weight:bold;">¥{q_grand_total:,}</span> &nbsp;|&nbsp;
                <b>粗利額:</b> <span style="color:#188038; font-size:16px; font-weight:bold;">¥{q_tot_profit:,}</span> (粗利率: <b>{q_tot_margin:.1f}%</b>)
            </div>
            """, unsafe_allow_html=True)

            row_cust_master = cust_lookup.get(q_cust_comp, {})
            cust_id_val = str(row_cust_master.get("加盟店ID", row_cust_master.get("顧客ID", -1))).replace(".0", "") if isinstance(row_cust_master, pd.Series) else -1
            cust_sei_val = str(row_cust_master.get("責任者姓", "ご担当者")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("責任者姓")) else "ご担当者"
            cust_mei_val = str(row_cust_master.get("責任者名", "様")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("責任者名")) else "様"
            cust_k_sei = str(row_cust_master.get("責任者姓カナ", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("責任者姓カナ")) else ""
            cust_k_mei = str(row_cust_master.get("責任者名カナ", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("責任者名カナ")) else ""
            cust_email_val = str(row_cust_master.get("メールアドレス", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("メールアドレス")) else ""
            cust_fax_val = str(row_cust_master.get("FAX番号", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("FAX番号")) else ""
            cust_pref_val = str(row_cust_master.get("都道府県", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("都道府県")) else ""
            cust_addr1_val = str(row_cust_master.get("住所１", q_cust_addr)).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("住所１")) else q_cust_addr
            cust_addr2_val = str(row_cust_master.get("住所２", "")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("住所２")) else ""
            cust_rank_label = str(row_cust_master.get("加盟店ランク", "ロット単価")).strip() if isinstance(row_cust_master, pd.Series) and pd.notna(row_cust_master.get("加盟店ランク")) else "ロット単価"

            current_estimate_obj = {
                "id": len(st.session_state.saved_estimates) + 840,
                "q_no": q_no,
                "staff": est_staff,
                "created_at": datetime.now().strftime("%Y/%m/%d %H:%M"),
                "q_date": q_date.strftime("%Y/%m/%d"),
                "q_exp_date": q_exp_date.strftime("%Y/%m/%d"),
                "payment": q_payment,
                "cust_id": cust_id_val,
                "cust_comp": q_cust_comp,
                "cust_person": q_cust_person,
                "cust_sei": cust_sei_val,
                "cust_mei": cust_mei_val,
                "cust_kana_sei": cust_k_sei,
                "cust_kana_mei": cust_k_mei,
                "cust_post": q_cust_post,
                "cust_pref": cust_pref_val,
                "cust_addr": q_cust_addr,
                "ship_comp": ship_cust_comp,
                "ship_person": ship_cust_person,
                "ship_sei": cust_sei_val,
                "ship_mei": cust_mei_val,
                "ship_post": ship_cust_post,
                "ship_pref": cust_pref_val,
                "ship_addr": ship_cust_addr,
                "ship_addr1": ship_cust_addr,
                "ship_addr2": "",
                "ship_tel": ship_cust_tel,
                "shipping_fee": q_shipping_fee,
                "cod_fee": 0,
                "sub_total": q_tot_sales,
                "tax": q_tax,
                "grand_total": q_grand_total,
                "memo": q_memo,
                "hitokoto": q_staff_memo,
                "items": updated_records
            }

            def generate_quote_excel_exact_style(df_calc):
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "見積書"
                ws.views.sheetView[0].showGridLines = True

                font_title = Font(name="Yu Gothic", size=18, bold=True)
                font_section = Font(name="Yu Gothic", size=11, bold=True)
                font_main = Font(name="Yu Gothic", size=9)
                font_bold = Font(name="Yu Gothic", size=9, bold=True)

                thin = Side(style='thin', color='A0A0A0')
                thick = Side(style='medium', color='000000')
                border_cell = Border(left=thin, right=thin, top=thin, bottom=thin)
                border_header = Border(left=thin, right=thin, top=thick, bottom=thick)
                fill_gray = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")

                ws.merge_cells("A2:F2")
                t_cell = ws["A2"]
                t_cell.value = "見 積 書"
                t_cell.font = font_title
                t_cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[2].height = 30
                for col_i in range(1, 7): ws.cell(row=3, column=col_i).border = Border(bottom=thick)

                ws["A5"] = f"〒 {q_cust_post}" if q_cust_post else ""
                ws["A6"] = q_cust_addr
                ws["A7"] = q_cust_comp
                ws["A8"] = q_cust_person
                ws["A5"].font = font_main; ws["A6"].font = font_main
                ws["A7"].font = font_section; ws["A8"].font = font_section

                ws["D7"] = "スマイルキッズ株式会社"; ws["D7"].font = font_bold
                ws["D8"] = "スマイルキッズ法人卸サイト"; ws["D8"].font = font_main
                ws["D9"] = "住所: 〒537-0003 大阪市東成区神路4丁目3番18号"; ws["D9"].font = font_main
                ws["D10"] = "TEL: 06-6976-1371  FAX: 06-6976-8940"; ws["D10"].font = font_main
                ws["D11"] = "Mail: smilekids-company@smile-asahi.co.jp"; ws["D11"].font = font_main
                ws["D12"] = "登録番号: T8120001247513"; ws["D12"].font = font_main

                if hanko_bytes:
                    try:
                        img_stream = io.BytesIO(hanko_bytes)
                        img = OpenpyxlImage(img_stream)
                        img.width = 65; img.height = 65
                        ws.add_image(img, "F7")
                    except Exception: pass

                ws.merge_cells("A14:F14")
                c_box1 = ws["A14"]
                c_box1.value = " 見積内容"
                c_box1.font = font_section; c_box1.fill = fill_gray; c_box1.border = border_cell
                ws.row_dimensions[14].height = 22

                ws["A15"] = "■見積金額"; ws["B15"] = f":  {q_grand_total:,} 円 (税込)"
                ws["A15"].font = font_bold; ws["B15"].font = font_bold
                ws["A16"] = "■お支払い方法"; ws["B16"] = f":  {q_payment}"
                ws["A16"].font = font_bold; ws["B16"].font = font_main

                ws.merge_cells("A18:F18")
                c_box2 = ws["A18"]
                c_box2.value = " お見積り明細"
                c_box2.font = font_section; c_box2.fill = fill_gray; c_box2.border = border_cell
                ws.row_dimensions[18].height = 22

                ws["A19"] = "■お届け先:"; ws["B19"] = f"〒 {ship_cust_post}" if ship_cust_post else ""
                ws["B20"] = ship_cust_addr; ws["B21"] = ship_cust_comp; ws["B22"] = ship_cust_person
                ws["B23"] = f"TEL: {ship_cust_tel}" if ship_cust_tel else ""
                ws["A19"].font = font_bold
                for r_i in range(19, 24): ws.cell(row=r_i, column=2).font = font_main

                ws["D19"] = "■お見積日"; ws["E19"] = f": {q_date.strftime('%Y/%m/%d')}"
                ws["D20"] = "■見積番号"; ws["E20"] = f": {q_no}"
                ws["D21"] = "■見積有効期限"; ws["E21"] = f": {q_exp_date.strftime('%Y/%m/%d')}"
                for r_i in range(19, 22):
                    ws.cell(row=r_i, column=4).font = font_bold
                    ws.cell(row=r_i, column=5).font = font_main

                start_row = 25
                headers = ["No.", "品名・規格", "JANコード", "数量", "単価", "金額"]
                ws.row_dimensions[start_row].height = 24
                for c_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=start_row, column=c_idx, value=h)
                    cell.font = font_bold; cell.fill = fill_gray; cell.border = border_header
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                cur_r = start_row + 1
                for idx_c, r in df_calc.iterrows():
                    ws.row_dimensions[cur_r].height = 20
                    ws.cell(row=cur_r, column=1, value=idx_c + 1).alignment = Alignment(horizontal="center")
                    ws.cell(row=cur_r, column=2, value=f"{r['商品名']} ({r['商品コード']})").alignment = Alignment(horizontal="left")
                    ws.cell(row=cur_r, column=3, value=str(r['JANコード'])).alignment = Alignment(horizontal="center")
                    ws.cell(row=cur_r, column=4, value=r['数量']).number_format = '#,##0'
                    ws.cell(row=cur_r, column=5, value=r['見積単価']).number_format = '#,##0'
                    ws.cell(row=cur_r, column=6, value=r['合計金額']).number_format = '#,##0'
                    for c_i in range(1, 7):
                        c = ws.cell(row=cur_r, column=c_i); c.font = font_main; c.border = border_cell
                    cur_r += 1

                ws.row_dimensions[cur_r].height = 20
                ws.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=3)
                tot_sub_lbl = ws.cell(row=cur_r, column=1, value="小計 / 合計")
                tot_sub_lbl.alignment = Alignment(horizontal="center", vertical="center")
                tot_sub_lbl.font = font_bold; tot_sub_lbl.fill = fill_gray

                c_tot_qty = ws.cell(row=cur_r, column=4, value=q_tot_qty); c_tot_qty.number_format = '#,##0'; c_tot_qty.font = font_bold; c_tot_qty.fill = fill_gray
                c_tot_p = ws.cell(row=cur_r, column=5, value=q_tot_unit_price); c_tot_p.number_format = '#,##0'; c_tot_p.font = font_bold; c_tot_p.fill = fill_gray
                c_tot_a = ws.cell(row=cur_r, column=6, value=q_tot_sales); c_tot_a.number_format = '#,##0'; c_tot_a.font = font_bold; c_tot_a.fill = fill_gray
                for c_i in range(1, 7): ws.cell(row=cur_r, column=c_i).border = border_cell
                cur_r += 1

                for s_lbl, s_val in [("税抜金額合計", q_tot_sales), ("送料 (税抜)", q_shipping_fee), ("消費税 (10%)", q_tax), ("税込合計金額", q_grand_total)]:
                    ws.row_dimensions[cur_r].height = 20
                    ws.merge_cells(start_row=cur_r, start_column=4, end_row=cur_r, end_column=5)
                    lbl_c = ws.cell(row=cur_r, column=4, value=s_lbl)
                    lbl_c.alignment = Alignment(horizontal="center", vertical="center")
                    lbl_c.font = font_bold; lbl_c.fill = fill_gray; lbl_c.border = border_cell
                    ws.cell(row=cur_r, column=5).border = border_cell
                    val_c = ws.cell(row=cur_r, column=6, value=s_val)
                    val_c.number_format = '#,##0'; val_c.font = font_bold; val_c.border = border_cell
                    cur_r += 1

                ws.column_dimensions['A'].width = 6; ws.column_dimensions['B'].width = 38; ws.column_dimensions['C'].width = 16
                ws.column_dimensions['D'].width = 12; ws.column_dimensions['E'].width = 13; ws.column_dimensions['F'].width = 15

                out = io.BytesIO()
                wb.save(out)
                return out.getvalue()

            hanko_base64 = base64.b64encode(hanko_bytes).decode("utf-8") if hanko_bytes else ""
            pdf_rows_html = ""
            for idx, r in calc_df.iterrows():
                p_display = f"{r['商品名']} ({r['商品コード']})" if r['商品名'] else str(r['商品コード'])
                pdf_rows_html += f"""
                <tr>
                    <td style="text-align:center;">{idx + 1}</td>
                    <td style="text-align:left;">{p_display}</td>
                    <td style="text-align:center;">{r['JANコード']}</td>
                    <td style="text-align:right;">{r['数量']:,}</td>
                    <td style="text-align:right;">¥{r['見積単価']:,}</td>
                    <td style="text-align:right;">¥{r['合計金額']:,}</td>
                </tr>
                """

            hanko_img_tag = f'<img src="data:image/png;base64,{hanko_base64}" style="position:absolute; right:15px; top:5px; width:75px; height:75px; opacity:0.92;" />' if hanko_bytes else ''

            pdf_template_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>御見積書_{q_cust_comp}</title>
                <style>
                    @media print {{
                        @page {{ size: A4 portrait; margin: 12mm; }}
                        body {{ -webkit-print-color-adjust: exact; }}
                        .no-print {{ display: none !important; }}
                    }}
                    body {{
                        font-family: -apple-system, BlinkMacSystemFont, "Yu Gothic", "Meiryo", sans-serif;
                        color: #111827; padding: 10px; max-width: 800px; margin: 0 auto;
                    }}
                    .title {{ text-align: center; font-size: 26px; font-weight: bold; letter-spacing: 12px; padding-bottom: 6px; border-bottom: 2px solid #111; margin-bottom: 20px; }}
                    .header-flex {{ display: flex; justify-content: space-between; margin-bottom: 15px; }}
                    .cust-area {{ font-size: 14px; line-height: 1.6; width: 55%; }}
                    .cust-comp {{ font-size: 17px; font-weight: bold; margin-top: 4px; }}
                    .cust-person {{ font-size: 17px; font-weight: bold; margin-top: 2px; }}
                    .company-area {{ position: relative; font-size: 12px; line-height: 1.5; width: 45%; padding-left: 20px; }}
                    .company-name {{ font-size: 15px; font-weight: bold; }}
                    .section-header {{ background: #e5e7eb; border: 1px solid #9ca3af; padding: 4px 10px; font-weight: bold; font-size: 14px; margin-top: 15px; margin-bottom: 10px; }}
                    .content-box {{ font-size: 14px; line-height: 1.8; margin-left: 10px; }}
                    .quote-table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 13px; }}
                    .quote-table th {{ background: #f3f4f6; border: 1px solid #9ca3af; padding: 7px 6px; }}
                    .quote-table td {{ border: 1px solid #d1d5db; padding: 7px 8px; font-variant-numeric: tabular-nums; }}
                    .summary-table {{ width: 340px; margin-left: auto; margin-top: 10px; border-collapse: collapse; font-size: 13px; }}
                    .summary-table td {{ border: 1px solid #9ca3af; padding: 6px 10px; }}
                    .btn-print {{ background: #2563eb; color: #fff; padding: 10px 22px; border: none; border-radius: 6px; font-size: 15px; font-weight: bold; cursor: pointer; }}
                </style>
            </head>
            <body>
                <div class="no-print" style="text-align:right; margin-bottom:12px;">
                    <button class="btn-print" onclick="window.print();">🖨️ PDFとして保存 / 印刷する</button>
                </div>
                <div class="title">見積書</div>
                <div class="header-flex">
                    <div class="cust-area">
                        <div>{'〒 ' + q_cust_post if q_cust_post else ''}</div>
                        <div>{q_cust_addr}</div>
                        <div class="cust-comp">{q_cust_comp}</div>
                        <div class="cust-person">{q_cust_person}</div>
                    </div>
                    <div class="company-area">
                        {hanko_img_tag}
                        <div class="company-name">スマイルキッズ株式会社</div>
                        <div>住所: 〒537-0003 大阪市東成区神路4丁目3番18号</div>
                        <div>TEL : 06-6976-1371 &nbsp; FAX : 06-6976-8940</div>
                    </div>
                </div>
                <div class="section-header">見積内容</div>
                <div class="content-box">
                    <div>■見積金額 &nbsp;&nbsp;&nbsp;&nbsp;: &nbsp;<b>{q_grand_total:,} 円</b> (税込)</div>
                    <div>■お支払い方法 : &nbsp;{q_payment}</div>
                </div>
                <table class="quote-table">
                    <thead>
                        <tr>
                            <th style="width: 35px;">No.</th>
                            <th>品名・規格</th>
                            <th style="width: 140px;">JANコード</th>
                            <th style="width: 75px;">数量</th>
                            <th style="width: 90px;">単価</th>
                            <th style="width: 105px;">金額</th>
                        </tr>
                    </thead>
                    <tbody>
                        {pdf_rows_html}
                        <tr style="background:#f9fafb; font-weight:bold; border-top:2px solid #9ca3af;">
                            <td colspan="3" style="text-align:center;">小計 / 合計</td>
                            <td style="text-align:right;">{q_tot_qty:,}</td>
                            <td style="text-align:right;">¥{q_tot_unit_price:,}</td>
                            <td style="text-align:right;">¥{q_tot_sales:,}</td>
                        </tr>
                    </tbody>
                </table>
                <table class="summary-table">
                    <tr>
                        <td style="background:#f3f4f6; text-align:center; font-weight:bold;">税抜金額合計</td>
                        <td style="text-align:right; font-weight:bold;">¥{q_tot_sales:,}</td>
                    </tr>
                    <tr>
                        <td style="background:#f3f4f6; text-align:center; font-weight:bold;">送料 (税抜)</td>
                        <td style="text-align:right; font-weight:bold;">¥{q_shipping_fee:,}</td>
                    </tr>
                    <tr>
                        <td style="background:#f3f4f6; text-align:center; font-weight:bold;">消費税 (10%)</td>
                        <td style="text-align:right; font-weight:bold;">¥{q_tax:,}</td>
                    </tr>
                    <tr style="background:#f9fafb;">
                        <td style="background:#e5e7eb; text-align:center; font-weight:bold; font-size:14px;">税込合計金額</td>
                        <td style="text-align:right; font-weight:bold; font-size:15px; color:#d93025;">¥{q_grand_total:,}</td>
                    </tr>
                </table>
            </body>
            </html>
            """

            file_prefix = q_cust_comp.replace(' ', '_') if q_cust_comp else "御見積書"
            btn_c1, btn_c2, btn_c3 = st.columns([1.4, 1.4, 1.6])
            with btn_c1:
                if st.button("💾 この見積を保存して一覧へ登録", use_container_width=True, key=f"save_btn_{s_key}"):
                    st.session_state.saved_estimates.insert(0, current_estimate_obj)
                    save_quotes_list(st.session_state.saved_estimates)
                    st.success("✅ 保存済み見積一覧に登録しました！")
            with btn_c2:
                st.download_button(
                    label="📥 御見積書（Excel・角印付き）",
                    data=generate_quote_excel_exact_style(calc_df),
                    file_name=f"御見積書_{file_prefix}_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key=f"dl_excel_{s_key}"
                )
            with btn_c3:
                st.download_button(
                    label="🛒 注文確定：スマレジ取込用CSV出力",
                    data=create_smaregi_order_csv_bytes(current_estimate_obj, raw_prod_title_map, jan_std_price_dict),
                    file_name=f"order_{q_no}_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                    mime="text/csv",
                    use_container_width=True,
                    key=f"dl_csv_{s_key}"
                )

            st.write("")
            st.markdown("##### 👁️ PDFプレビュー表示＆ワンクリックPDF保存")
            components.html(pdf_template_html, height=720, scrolling=True)

    elif tab_choice.startswith("📋"):
        st.markdown(f"##### 📋 保存済み見積一覧（担当: {est_staff}）")
        if not staff_Filter_list:
            st.info(f"担当「{est_staff}」の保存された見積もりはありません。")
        else:
            search_est_kw = st.text_input("🔍 見積番号・会社名・商品名で絞り込み", placeholder="例: アーネスト, ABA-105...")
            for idx_q, est in enumerate(st.session_state.saved_estimates):
                if est_staff != "（すべて表示）" and est.get("staff", "") != est_staff and (est_staff not in str(est.get("hitokoto", ""))):
                    continue
                if search_est_kw.strip():
                    if search_est_kw.strip() not in json.dumps(est, ensure_ascii=False):
                        continue

                st.markdown(f"""
                <div class="est-card">
                    <div class="est-header">
                        <div>
                            <span class="badge-green">店</span> &nbsp;
                            <b>見積番号: {est.get('q_no', '')}</b> &nbsp;|&nbsp;
                            <b>担当:</b> {est.get('staff', est_staff)} &nbsp;|&nbsp;
                            <b>見積日時:</b> {est.get('created_at', '')}
                        </div>
                        <div>
                            <b>税込見積総額:</b> <span style="color:#d93025; font-size:16px; font-weight:bold;">¥{int(est.get('grand_total', 0)):,}</span>
                        </div>
                    </div>
                </div>
                """, unsafe_allow_html=True)

                c_info1, c_info2, c_items = st.columns([1.5, 1.5, 2.5])
                with c_info1:
                    st.markdown("**【宛先情報】**")
                    st.write(f"🏢 **{est.get('cust_comp', '')}**")
                    st.write(f"👤 {est.get('cust_person', '')}")
                    st.write(f"〒{est.get('cust_post', '')} {est.get('cust_addr', '')}")
                with c_info2:
                    st.markdown("**【お届け先・備考】**")
                    st.write(f"🏢 {est.get('ship_comp', '')}")
                    if est.get('memo'): st.info(f"伝達事項: {est.get('memo')}")
                with c_items:
                    items_df = pd.DataFrame(est.get("items", []))
                    if not items_df.empty:
                        show_sub = items_df[[c for c in ["商品コード", "商品名", "数量", "見積単価", "合計金額"] if c in items_df.columns]].copy()
                        st.dataframe(show_sub, use_container_width=True, hide_index=True)

                b_col1, b_col2, b_col3, _ = st.columns([2.0, 1.6, 1.0, 2.0])
                with b_col1:
                    st.download_button(
                        label="🛒 注文確定：スマレジ用CSV出力",
                        data=create_smaregi_order_csv_bytes(est, raw_prod_title_map, jan_std_price_dict),
                        file_name=f"order_{est.get('q_no')}_{datetime.now().strftime('%Y%m%d')}.csv",
                        mime="text/csv",
                        key=f"dl_saved_csv_{idx_q}_{est.get('q_no')}"
                    )
                with b_col2:
                    if st.button("✏️ 編集画面に読込", key=f"load_est_{idx_q}"):
                        st.session_state["load_target_estimate"] = est
                        st.rerun()
                with b_col3:
                    if st.button("🗑️ 削除", key=f"del_est_{idx_q}"):
                        st.session_state.saved_estimates.pop(idx_q)
                        save_quotes_list(st.session_state.saved_estimates)
                        st.rerun()
                st.markdown("<hr style='margin:12px 0; border:0; border-top:1px dashed #cbd5e1;'>", unsafe_allow_html=True)

    st.stop()

# ========================================================
# ★ 🏷️ ランク価格マスタ一覧（指定12ランク限定・ロット数列削除）
# ========================================================
if view_type == '🏷️ ランク価格マスタ一覧':
    raw_rank_path, raw_rank_mtime = get_rank_price_file_info()
    raw_rank_df = load_rank_price_master(raw_rank_path, raw_rank_mtime)

    if raw_rank_df.empty and member_p_path and os.path.exists(member_p_path):
        raw_rank_df = load_rank_price_master(member_p_path, member_p_mtime)

    if raw_rank_df.empty:
        st.error("ランク価格データが見つかりませんでした。")
        st.stop()

    code_to_name_map = df.dropna(subset=["商品コード", "商品名"]).drop_duplicates(subset=["商品コード"]).set_index("商品コード")["商品名"].to_dict()
    for k_t, v_t in jan_title_dict.items():
        if str(k_t).strip() not in code_to_name_map:
            code_to_name_map[str(k_t).strip()] = str(v_t).strip()

    rank_col = next((c for c in raw_rank_df.columns if "ランク名" in str(c)), None)
    if rank_col is None:
        rank_col = next((c for c in raw_rank_df.columns if any(k in str(c) for k in ["ランク", "会社名", "取引先", "グループ"])), None)
    sku_col = next((c for c in raw_rank_df.columns if any(k in str(c) for k in ["商品番号", "商品コード", "商品管理番号", "SKU", "品番"])), raw_rank_df.columns[0])
    price_col = next((c for c in raw_rank_df.columns if any(k in str(c) for k in ["ロット単価", "単価", "下代", "価格"])), None)
    cost_col = next((c for c in raw_rank_df.columns if "原価" in str(c)), None)

    work_rk = pd.DataFrame()
    work_rk["商品番号"] = raw_rank_df[sku_col].astype(str).str.strip()
    work_rk = work_rk[~work_rk["商品番号"].isin(["", "nan", "None"])].copy()

    if "商品名" in raw_rank_df.columns:
        work_rk["商品名"] = raw_rank_df["商品名"].astype(str).str.strip()
    else:
        work_rk["商品名"] = work_rk["商品番号"].map(lambda c: code_to_name_map.get(c, jan_title_dict.get(c, "")))

    if rank_col is not None:
        work_rk["ランク名"] = raw_rank_df[rank_col].astype(str).str.strip()
    else:
        work_rk["ランク名"] = "1.1倍"

    # 🔴 指定された12種類のみに厳密に絞り込み
    work_rk = work_rk[work_rk["ランク名"].isin(ALLOWED_EXACT_RANKS)].copy()

    work_rk["ロット単価"] = raw_rank_df[price_col].apply(clean_num).astype(int) if price_col is not None else 0
    work_rk["原価"] = raw_rank_df[cost_col].apply(clean_num).astype(int) if cost_col is not None else work_rk["商品番号"].map(lambda c: int(cost_master_dict.get(c, 0)))
    work_rk = work_rk[(work_rk["ロット単価"] > 0) & (work_rk["ロット単価"] < 90000)].copy()

    r_col1, r_col2, r_col3 = st.columns([1.5, 1.5, 1.0])
    with r_col1:
        # 指定された12種類順でプルダウン表示
        avail_ranks = [rk for rk in ALLOWED_EXACT_RANKS if rk in work_rk["ランク名"].values]
        sel_rank = st.selectbox("🎯 ランクで絞り込み", ["（すべて）"] + avail_ranks)
    with r_col2:
        search_rank_prod = st.text_input("🔍 商品番号・商品名で検索", placeholder="例：ABO-03...")
    with r_col3:
        rank_limit_val = st.selectbox("表示件数", [50, 100, 200, 500, "全件"], index=0)

    filt_rank_df = work_rk.copy()
    if sel_rank != "（すべて）":
        filt_rank_df = filt_rank_df[filt_rank_df["ランク名"] == sel_rank]
    if search_rank_prod.strip():
        kw_rp = search_rank_prod.strip()
        filt_rank_df = filt_rank_df[filt_rank_df["商品番号"].str.contains(kw_rp, case=False, na=False) | filt_rank_df["商品名"].str.contains(kw_rp, case=False, na=False)]

    filt_rank_df["ロット粗利額"] = filt_rank_df["ロット単価"] - filt_rank_df["原価"]
    filt_rank_df["粗利率"] = filt_rank_df.apply(lambda r: f"{(r['ロット粗利額'] / r['ロット単価'] * 100):.1f}%" if r['ロット単価'] > 0 else "-", axis=1)

    # 🔴 ロット数 列を完全に削除
    disp_rank_cols = ["商品番号", "商品名", "ランク名", "ロット単価", "原価", "ロット粗利額", "粗利率"]
    out_rank_df = filt_rank_df[disp_rank_cols].reset_index(drop=True)
    out_rank_df.insert(0, "番号", range(1, len(out_rank_df) + 1))

    r_top_c1, r_top_c2 = st.columns([4, 1.5])
    with r_top_c1:
        st.markdown(f"""
        <div class="total-card" style="margin-bottom:0;">
            <b>対象ランク:</b> <span style="color:#1a73e8; font-weight:bold;">{sel_rank}</span> &nbsp;|&nbsp;
            <b>該当件数:</b> <span style="font-size:16px; font-weight:bold; color:#d93025;">{len(out_rank_df):,} 件</span>
        </div>
        """, unsafe_allow_html=True)
    with r_top_c2:
        st.download_button(
            label="📥 ランク価格一覧をExcel保存",
            data=generate_excel_bytes(out_rank_df, "ランク価格一覧"),
            file_name=f"ランク価格_{sel_rank}_{datetime.now().strftime('%Y%m%d')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

    st.write("")
    if rank_limit_val != "全件":
        out_rank_df = out_rank_df.head(int(rank_limit_val))
    render_universal_table(out_rank_df)
    st.stop()

# ========================================================
# ★ 期間設定（売上分析用）
# ========================================================
idx = valid_ym.index(latest_ym) if latest_ym in valid_ym else len(valid_ym) - 1
curr_y = int(latest_ym[:4])
curr_m = int(latest_ym[5:7])
fiscal_year = curr_y if curr_m >= 7 else curr_y - 1

if quick_choice == '当月':
    s_def, e_def = latest_ym, latest_ym
    s_date_calc = datetime(curr_y, curr_m, 1).date()
    e_date_calc = max_date
elif quick_choice == '前月':
    prev = valid_ym[max(0, idx - 1)]
    s_def, e_def = prev, prev
    py, pm = int(prev[:4]), int(prev[5:7])
    s_date_calc = datetime(py, pm, 1).date()
    e_date_calc = (pd.Timestamp(year=py, month=pm, day=1) + pd.offsets.MonthEnd(1)).date()
elif quick_choice == '当期累計':
    s_def, e_def = f"{fiscal_year}年07月", latest_ym
    s_date_calc = datetime(fiscal_year, 7, 1).date()
    e_date_calc = max_date
elif quick_choice == '前期累計':
    s_def, e_def = f"{fiscal_year - 1}年07月", f"{fiscal_year}年06月"
    s_date_calc = datetime(fiscal_year - 1, 7, 1).date()
    e_date_calc = datetime(fiscal_year, 6, 30).date()
elif quick_choice == '全データ累計':
    s_def, e_def = valid_ym[0], valid_ym[-1]
    s_date_calc = min_date
    e_date_calc = max_date
else:
    s_def, e_def = valid_ym[0], latest_ym
    s_date_calc = min_date
    e_date_calc = max_date

if "prev_quick_choice" not in st.session_state or st.session_state.prev_quick_choice != quick_choice:
    st.session_state.prev_quick_choice = quick_choice
    if quick_choice != "手動指定":
        sync_ym = s_def if s_def in desc_valid_ym else desc_valid_ym[0]
        st.session_state["analysis_ym_start_v3"] = sync_ym
        st.session_state["analysis_ym_end_v3"] = e_def if e_def in desc_valid_ym else desc_valid_ym[0]
        st.session_state["tgt_ym_meeting_v3"] = sync_ym
        st.session_state["tgt_ym_topcust_v3"] = sync_ym

if "analysis_ym_start_v3" not in st.session_state:
    st.session_state["analysis_ym_start_v3"] = s_def if s_def in desc_valid_ym else desc_valid_ym[0]
if "analysis_ym_end_v3" not in st.session_state:
    st.session_state["analysis_ym_end_v3"] = e_def if e_def in desc_valid_ym else desc_valid_ym[0]

if view_type not in ['📑 営業会議用：上位客先別売上実績', '📊 営業会議用：上位社 売上比較シート']:
    col_mode, col_p1, col_p2, col_staff = st.columns([1, 1.2, 1.2, 1.6])
    with col_mode:
        date_mode = st.radio("指定単位", ["年月で指定", "日で指定"], horizontal=True, key="date_mode_radio_analysis")

    if date_mode == "年月で指定":
        with col_p1: ym_start = st.selectbox("開始年月", desc_valid_ym, key="analysis_ym_start_v3")
        with col_p2: ym_end = st.selectbox("終了年月", desc_valid_ym, key="analysis_ym_end_v3")
        st.session_state["tgt_ym_meeting_v3"] = ym_start
        st.session_state["tgt_ym_topcust_v3"] = ym_start
    else:
        with col_p1: d_start = st.date_input("開始日", value=s_date_calc, min_value=min_date, max_value=max_date)
        with col_p2: d_end = st.date_input("終了日", value=e_date_calc, min_value=min_date, max_value=max_date)

    with col_staff:
        staff_options = ["（すべて）"] + sorted([x for x in df["担当者"].unique() if x and x != "nan"])
        sel_staff = st.selectbox("担当者指定", staff_options)

    all_unique_custs = sorted([x for x in df["得意先名"].unique() if x])

    c1, c2, c3, c4 = st.columns(4)
    with c1: sel_cust = st.selectbox("得意先名", ["（すべて）"] + all_unique_custs)
    with c2: search_cust = st.text_input("得意先検索", placeholder="会社名の一部を入力（例：丸藤、アマゾン）...")
    with c3:
        prod_pairs = df[["商品コード", "商品名"]].dropna().drop_duplicates()
        prod_list = ["（すべて）"] + sorted([f"{r['商品コード']} : {r['商品名']}" for _, r in prod_pairs.iterrows()])
        sel_prod = st.selectbox("商品選択", prod_list)
    with c4: search_prod = st.text_input("商品検索", placeholder="商品コード・品名の一部を入力...")

    matched_custs = []
    chosen_sub_branch = "（すべて合算）"
    if search_cust.strip():
        kw_clean = re.sub(r'[\s ]+', '', search_cust.strip())
        matched_custs = [c for c in all_unique_custs if kw_clean in re.sub(r'[\s ]+', '', c)]

    if len(matched_custs) > 1:
        st.info(f"💡 「{search_cust.strip()}」で **{len(matched_custs)} 件** の得意先が見つかりました：")
        chosen_sub_branch = st.selectbox("🎯 対象の取引先を選択", [f"（該当する全{len(matched_custs)}件を合算表示）"] + matched_custs)

    limit_col, _, _, _ = st.columns(4)
    with limit_col: limit_val = st.selectbox("表示件数", [50, 100, 200, 500, "全件"], index=0)
else:
    date_mode = "年月で指定"
    ym_start = st.session_state.get("analysis_ym_start_v3", latest_ym)
    ym_end = st.session_state.get("analysis_ym_end_v3", latest_ym)
    sel_staff = "（すべて）"
    sel_cust = "（すべて）"; search_cust = ""; sel_prod = "（すべて）"; search_prod = ""; limit_val = 50; matched_custs = []; chosen_sub_branch = ""

prod_first_seen_dt = df.groupby("商品コード")["日付_dt"].min().to_dict()
prod_last_seen_dt = df.groupby("商品コード")["日付_dt"].max().to_dict()
code_to_name = df.dropna(subset=["商品コード", "商品名"]).drop_duplicates(subset=["商品コード"]).set_index("商品コード")["商品名"].to_dict()

one_year_ago_date = pd.Timestamp(max_date) - pd.DateOffset(years=1)

if view_type in ['📑 営業会議用：上位客先別売上実績', '📊 営業会議用：上位社 売上比較シート']:
    filtered = df.copy(); ly_filtered = pd.DataFrame(); period_text = ""
elif view_type == '💤 滞留・未販売商品分析（売れてない商品）':
    filtered = df.copy(); period_text = f"基準日: {max_date.strftime('%Y/%m/%d')}"; ly_filtered = pd.DataFrame()
elif view_type == '✨ 新商品売上分析（直近1年）':
    true_new_skus = {sku for sku, f_dt in prod_first_seen_dt.items() if pd.notna(f_dt) and f_dt >= one_year_ago_date}
    filtered = df[(df["日付_dt"] >= one_year_ago_date) & (df["商品コード"].isin(true_new_skus))].copy()
    period_text = f"{one_year_ago_date.strftime('%Y/%m/%d')} 〜 {max_date.strftime('%Y/%m/%d')}（直近1年間に初登場した新商品）"
    ly_filtered = pd.DataFrame()
else:
    if date_mode == "年月で指定":
        s_ym_calc, e_ym_calc = (ym_end, ym_start) if ym_start > ym_end else (ym_start, ym_end)
        filtered = df[(df["年月"] >= s_ym_calc) & (df["年月"] <= e_ym_calc)]
        period_text = f"{s_ym_calc} 〜 {e_ym_calc}" if s_ym_calc != e_ym_calc else f"{s_ym_calc}（単月）"
        s_y, s_m = int(s_ym_calc[:4]), int(s_ym_calc[5:7])
        e_y, e_m = int(e_ym_calc[:4]), int(e_ym_calc[5:7])
        ly_filtered = df[(df["年月"] >= f"{s_y - 1}年{s_m:02d}月") & (df["年月"] <= f"{e_y - 1}年{e_m:02d}月")]
    else:
        dt_s, dt_e = pd.to_datetime(d_start), pd.to_datetime(d_end)
        if dt_s > dt_e: dt_s, dt_e = dt_e, dt_s
        filtered = df[(df["日付_dt"] >= dt_s) & (df["日付_dt"] <= dt_e)]
        period_text = f"{dt_s.strftime('%Y/%m/%d')} 〜 {dt_e.strftime('%Y/%m/%d')}"
        ly_filtered = df[(df["日付_dt"] >= (dt_s - pd.DateOffset(years=1))) & (df["日付_dt"] <= (dt_e - pd.DateOffset(years=1)))]

if view_type not in ['📑 営業会議用：上位客先別売上実績', '📊 営業会議用：上位社 売上比較シート']:
    if sel_cust != "（すべて）":
        filtered = filtered[filtered["得意先名"] == sel_cust]
        if not ly_filtered.empty: ly_filtered = ly_filtered[ly_filtered["得意先名"] == sel_cust]
    if search_cust.strip():
        kw_clean = re.sub(r'[\s ]+', '', search_cust.strip())
        if len(matched_custs) > 1 and not chosen_sub_branch.startswith("（該当する全"):
            filtered = filtered[filtered["得意先名"] == chosen_sub_branch]
            if not ly_filtered.empty: ly_filtered = ly_filtered[ly_filtered["得意先名"] == chosen_sub_branch]
        else:
            filtered = filtered[filtered["得意先名"].map(lambda x: kw_clean in re.sub(r'[\s ]+', '', str(x)))]
            if not ly_filtered.empty: ly_filtered = ly_filtered[ly_filtered["得意先名"].map(lambda x: kw_clean in re.sub(r'[\s ]+', '', str(x)))]

    if sel_staff != "（すべて）":
        filtered = filtered[filtered["担当者"] == sel_staff]
        if not ly_filtered.empty: ly_filtered = ly_filtered[ly_filtered["担当者"] == sel_staff]

    if sel_prod != "（すべて）":
        code = sel_prod.split(" : ")[0]
        filtered = filtered[filtered["商品コード"] == code]
        if not ly_filtered.empty: ly_filtered = ly_filtered[ly_filtered["商品コード"] == code]
    if search_prod.strip():
        kw_p = search_prod.strip()
        cond = filtered["商品名"].str.contains(kw_p, na=False) | filtered["商品コード"].str.contains(kw_p, na=False)
        filtered = filtered[cond]
        if not ly_filtered.empty:
            ly_cond = ly_filtered["商品名"].str.contains(kw_p, na=False) | ly_filtered["商品コード"].str.contains(kw_p, na=False)
            ly_filtered = ly_filtered[ly_cond]

    total_sales = int(filtered["金額"].sum()) if not filtered.empty else 0

# ========================================================
# ★ 🏢 得意先別ランキング（行クリックでTOP10詳細POP完全復元）
# ========================================================
if view_type == '🏢 得意先別ランキング':
    group_cols = [c for c in ["得意先名", "担当者"] if c in filtered.columns]
    agg = filtered.groupby(group_cols).agg(販売金額=("金額", "sum"), 数量=("数量", "sum"), 原価合計=("原価合計", "sum"), 粗利額=("粗利額", "sum"), 取扱品目数=("商品コード", "nunique")).reset_index() if not filtered.empty else pd.DataFrame()
    if not agg.empty:
        agg["粗利率"] = agg.apply(lambda r: f"{(r['粗利額'] / r['販売金額'] * 100):.1f}%" if r['販売金額'] > 0 else "-", axis=1)
        agg["構成比"] = (agg["販売金額"] / total_sales * 100).round(2).map(lambda x: f"{x:.2f}%") if total_sales > 0 else "0.00%"
        agg = agg.sort_values(by="販売金額", ascending=False)
        agg.insert(0, "順位", range(1, len(agg) + 1))
        if limit_val != "全件": agg = agg.head(int(limit_val))
        show_df = agg[["順位", "得意先名", "担当者", "販売金額", "粗利額", "粗利率", "原価合計", "数量", "取扱品目数", "構成比"]].copy()
    else:
        show_df = pd.DataFrame()

    c_sum1, c_sum2 = st.columns([4, 1.2])
    with c_sum1:
        st.markdown(f"""
        <div class="total-card" style="margin-bottom:0;">
            <b>抽出期間:</b> <span style="color:#1a73e8; font-weight:bold;">{period_text}</span> &nbsp;|&nbsp;
            <b>売上合計:</b> <span style="color:#d93025; font-size:16px; font-weight:bold;">{int(total_sales):,} 円</span>
        </div>
        """, unsafe_allow_html=True)
    with c_sum2:
        if not show_df.empty:
            st.download_button(
                label="📥 Excelで保存",
                data=generate_excel_bytes(show_df, "得意先別ランキング"),
                file_name=f"得意先別ランキング_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    st.write("")
    if not show_df.empty:
        st.info("👆 **表の行（得意先名）をクリックすると、その得意先の「売れ筋・売上TOP10詳細」がポップアップ表示されます。**")
        event = st.dataframe(
            show_df.style.format({"販売金額": "¥{:,}", "粗利額": "¥{:,}", "原価合計": "¥{:,}", "数量": "{:,}", "取扱品目数": "{:,}"}),
            use_container_width=True,
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            height=560
        )
        if event and event.selection and event.selection.rows:
            sel_row_idx = event.selection.rows[0]
            clicked_cust = show_df.iloc[sel_row_idx]["得意先名"]
            cust_sub_df = filtered[filtered["得意先名"] == clicked_cust]
            show_customer_top10_dialog(clicked_cust, cust_sub_df, period_text)
    else:
        st.info("該当期間のデータがありません。")
    st.stop()

elif view_type == '📦 商品別売上ランキング':
    group_cols = [c for c in ["得意先名", "担当者", "商品コード", "商品名"] if c in filtered.columns]
    if not filtered.empty:
        valid_date_df = filtered[filtered["日付"].str.strip() != ""]
        date_agg = valid_date_df.groupby(group_cols).agg(開始日=("日付", "min"), 終了日=("日付", "max")).reset_index()
        agg = filtered.groupby(group_cols).agg(販売金額=("金額", "sum"), 数量=("数量", "sum"), 単価原価=("単価原価", "first"), 原価合計=("原価合計", "sum"), 粗利額=("粗利額", "sum"), 単価_min=("単価", "min"), 単価_max=("単価", "max")).reset_index()
        agg = pd.merge(agg, date_agg, on=group_cols, how="left")
        agg["粗利率"] = agg.apply(lambda r: f"{(r['粗利額'] / r['販売金額'] * 100):.1f}%" if r['販売金額'] > 0 else "-", axis=1)
        agg["販売価格帯"] = agg.apply(lambda r: f"{int(r['単価_min']):,} 〜 {int(r['単価_max']):,}" if pd.notna(r['単価_min']) else "-", axis=1)
        agg["販売期間"] = agg.apply(lambda r: f"{r['開始日']} 〜 {r['終了日']}" if pd.notna(r['開始日']) else "-", axis=1)
        agg["構成比"] = (agg["販売金額"] / total_sales * 100).round(2).map(lambda x: f"{x:.2f}%") if total_sales > 0 else "0.00%"
        agg = agg.sort_values(by="販売金額", ascending=False)
        agg.insert(0, "番号", range(1, len(agg) + 1))
        if limit_val != "全件": agg = agg.head(int(limit_val))
        cols = ["番号", "得意先名", "担当者", "商品コード", "商品名", "数量", "販売価格帯", "単価原価", "販売金額", "原価合計", "粗利額", "粗利率", "構成比", "販売期間"]
        show_df = agg[[c for c in cols if c in agg.columns]].copy()
    else:
        show_df = pd.DataFrame()

    c_sum1, c_sum2 = st.columns([4, 1.2])
    with c_sum1:
        st.markdown(f"""
        <div class="total-card" style="margin-bottom:0;">
            <b>抽出期間:</b> <span style="color:#1a73e8; font-weight:bold;">{period_text}</span> &nbsp;|&nbsp;
            <b>販売金額合計:</b> <span style="color:#d93025; font-size:16px; font-weight:bold;">{int(total_sales):,} 円</span>
        </div>
        """, unsafe_allow_html=True)
    with c_sum2:
        if not show_df.empty:
            st.download_button(
                label="📥 Excelで保存",
                data=generate_excel_bytes(show_df, "商品別売上ランキング"),
                file_name=f"商品別売上ランキング_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    st.write("")
    render_universal_table(show_df)
    st.stop()

elif view_type == '👤 担当者別 売上＆目標達成分析':
    st.markdown(f"#### 👤 担当者別 売上＆目標達成分析（対象期間: {period_text}）")
    if not filtered.empty:
        staff_agg = filtered.groupby("担当者").agg(
            販売金額=("金額", "sum"),
            数量=("数量", "sum"),
            原価合計=("原価合計", "sum"),
            粗利額=("粗利額", "sum"),
            担当得意先数=("得意先名", "nunique"),
            取扱品目数=("商品コード", "nunique")
        ).reset_index()
        staff_agg = staff_agg[staff_agg["担当者"].astype(str).str.strip() != ""]

        if not ly_filtered.empty:
            ly_st = ly_filtered.groupby("担当者").agg(前年売上額=("金額", "sum"), 前年粗利額=("粗利額", "sum")).reset_index()
            staff_agg = pd.merge(staff_agg, ly_st, on="担当者", how="left").fillna(0)
        else:
            staff_agg["前年売上額"] = 0
            staff_agg["前年粗利額"] = 0

        staff_agg["前年売上額"] = staff_agg["前年売上額"].astype(int)
        staff_agg["増減額"] = staff_agg["販売金額"] - staff_agg["前年売上額"]
        staff_agg["前年比"] = staff_agg.apply(lambda r: f"{(r['販売金額'] / r['前年売上額'] * 100):.1f}%" if r["前年売上額"] > 0 else "-", axis=1)
        staff_agg["粗利率"] = staff_agg.apply(lambda r: f"{(r['粗利額'] / r['販売金額'] * 100):.1f}%" if r['販売金額'] > 0 else "-", axis=1)
        staff_agg["構成比"] = (staff_agg["販売金額"] / total_sales * 100).round(1).map(lambda x: f"{x:.1f}%") if total_sales > 0 else "0.0%"
        staff_agg = staff_agg.sort_values(by="販売金額", ascending=False)
        staff_agg.insert(0, "順位", range(1, len(staff_agg) + 1))

        cols_st = ["順位", "担当者", "販売金額", "前年売上額", "前年比", "増減額", "粗利額", "粗利率", "数量", "担当得意先数", "取扱品目数", "構成比"]
        show_st_df = staff_agg[[c for c in cols_st if c in staff_agg.columns]].copy()

        tot_prof = int(staff_agg["粗利額"].sum())
        tot_ly = int(staff_agg["前年売上額"].sum())
        yoy_tot_str = f"{(total_sales / tot_ly * 100):.1f}%" if tot_ly > 0 else "-"

        c_sum1, c_sum2 = st.columns([4, 1.2])
        with c_sum1:
            st.markdown(f"""
            <div class="total-card" style="margin-bottom:0;">
                <b>期間売上合計:</b> <span style="color:#d93025; font-size:16px; font-weight:bold;">¥{total_sales:,}</span> &nbsp;|&nbsp;
                <b>前年同期合計:</b> ¥{tot_ly:,} (前年比: <b>{yoy_tot_str}</b>) &nbsp;|&nbsp;
                <b>粗利合計:</b> <span style="color:#188038; font-weight:bold;">¥{tot_prof:,}</span>
            </div>
            """, unsafe_allow_html=True)
        with c_sum2:
            st.download_button(
                label="📥 Excelで保存",
                data=generate_excel_bytes(show_st_df, "担当者別分析"),
                file_name=f"担当者別売上分析_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        st.write("")
        render_universal_table(show_st_df)
    else:
        st.info("該当期間のデータがありません。")
    st.stop()

elif view_type == '✨ 新商品売上分析（直近1年）':
    st.markdown(f"#### ✨ 新商品売上分析（{period_text}）")
    st.caption("※全データ履歴の中で「初回売上日」が直近1年以内の新商品のみを自動抽出しています。")
    if not filtered.empty:
        c_sort1, c_sort2, _ = st.columns([1.8, 1.8, 2.0])
        with c_sort1:
            np_sort_col = st.radio(
                "📊 並び替え項目",
                ["販売金額", "数量", "導入社数"],
                horizontal=True,
                key="np_sort_col_radio"
            )
        with c_sort2:
            np_sort_order = st.radio(
                "🔃 並び順",
                ["⬇️ 降順（高い順）", "⬆️ 昇順（低い順）"],
                horizontal=True,
                key="np_sort_order_radio"
            )
        is_asc = np_sort_order.startswith("⬆️")

        np_agg = filtered.groupby(["商品コード", "商品名"]).agg(
            販売金額=("金額", "sum"),
            数量=("数量", "sum"),
            単価原価=("単価原価", "first"),
            原価合計=("原価合計", "sum"),
            粗利額=("粗利額", "sum"),
            導入社数=("得意先名", "nunique"),
            最終売上日=("日付", "max")
        ).reset_index()
        np_agg["初回売上日"] = np_agg["商品コード"].map(lambda c: prod_first_seen_dt.get(c).strftime('%Y/%m/%d') if pd.notna(prod_first_seen_dt.get(c)) else "-")
        np_agg["粗利率"] = np_agg.apply(lambda r: f"{(r['粗利額'] / r['販売金額'] * 100):.1f}%" if r['販売金額'] > 0 else "-", axis=1)
        np_agg = np_agg.sort_values(by=[np_sort_col, "販売金額"], ascending=[is_asc, False]).reset_index(drop=True)
        np_agg.insert(0, "順位", range(1, len(np_agg) + 1))
        if limit_val != "全件": np_agg = np_agg.head(int(limit_val))
        cols_np = ["順位", "商品コード", "商品名", "初回売上日", "最終売上日", "導入社数", "数量", "単価原価", "販売金額", "粗利額", "粗利率"]
        show_np_df = np_agg[[c for c in cols_np if c in np_agg.columns]].copy()

        c_sum1, c_sum2 = st.columns([4, 1.2])
        with c_sum1:
            st.markdown(f"""
            <div class="total-card" style="margin-bottom:0;">
                <b>新商品 該当品目数:</b> <span style="color:#1a73e8; font-weight:bold;">{len(show_np_df):,} 品目</span> &nbsp;|&nbsp;
                <b>新商品 売上合計:</b> <span style="color:#d93025; font-size:16px; font-weight:bold;">¥{int(show_np_df['販売金額'].sum()):,}</span> &nbsp;|&nbsp;
                <b>現在の並び順:</b> <span style="color:#188038; font-weight:bold;">{np_sort_col}（{'昇順' if is_asc else '降順'}）</span>
            </div>
            """, unsafe_allow_html=True)
        with c_sum2:
            st.download_button(
                label="📥 Excelで保存",
                data=generate_excel_bytes(show_np_df, "新商品売上分析"),
                file_name=f"新商品売上分析_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        st.write("")
        render_universal_table(show_np_df)
    else:
        st.info("直近1年間に初登場した新商品の売上データはありません。")
    st.stop()

elif view_type == '📊 営業会議用：上位社 売上比較シート':
    if "tgt_ym_meeting_v3" not in st.session_state:
        st.session_state["tgt_ym_meeting_v3"] = ym_start if ym_start in desc_valid_ym else desc_valid_ym[0]

    col_m1, col_m2, col_m3 = st.columns([1.5, 1.5, 1.5])
    with col_m1:
        target_meeting_ym = st.selectbox("対象月度（期間ショートカット連動）", desc_valid_ym, key="tgt_ym_meeting_v3")
    with col_m2:
        compare_cust_count = st.selectbox("比較上位社数", [4, 6, 8, 10, 20, 30], index=0, key="compare_cust_count_meeting_v3")
    with col_m3:
        staff_options_meeting = ["（すべて）"] + sorted([x for x in df["担当者"].unique() if x and x != "nan"])
        sel_staff_m = st.selectbox("担当者指定", staff_options_meeting, key="meeting_staff_sel_v3")

    tgt_y = int(target_meeting_ym[:4]); tgt_m = int(target_meeting_ym[5:7])
    ly_ym = f"{tgt_y - 1}年{tgt_m:02d}月"
    next_y = tgt_y + 1 if tgt_m == 12 else tgt_y
    next_m = 1 if tgt_m == 12 else tgt_m + 1
    next_ym = f"{next_y}年{next_m:02d}月"
    f_year = tgt_y if tgt_m >= 7 else tgt_y - 1
    prev_fy_ym = f"{next_y - 1}年{next_m:02d}月"

    df_scope = df if sel_staff_m == "（すべて）" else df[df["担当者"] == sel_staff_m]

    # 🔴 自己参照を完全に排除した安全な抽出
    cur_m_df = df_scope[df_scope["年月"] == target_meeting_ym] if not df_scope.empty else pd.DataFrame()
    ly_m_df = df_scope[df_scope["年月"] == ly_ym] if not df_scope.empty else pd.DataFrame()
    next_m_df = df_scope[df_scope["年月"] == next_ym] if not df_scope.empty else pd.DataFrame()
    prev_fy_df = df_scope[df_scope["年月"] == prev_fy_ym] if not df_scope.empty else pd.DataFrame()

    cust_id_lookup = {}
    if not customer_master_df.empty:
        c_name_col = "加盟店名" if "加盟店名" in customer_master_df.columns else customer_master_df.columns[0]
        c_id_col = "加盟店ID" if "加盟店ID" in customer_master_df.columns else ("顧客ID" if "顧客ID" in customer_master_df.columns else None)
        if c_id_col:
            for _, r_c in customer_master_df.iterrows():
                cn = clean_cust_name(str(r_c[c_name_col]))
                cid = str(r_c[c_id_col]).replace(".0", "").strip()
                if cn and cid != "nan": cust_id_lookup[cn] = cid

    cur_cust_sales = cur_m_df.groupby("得意先名")["金額"].sum() if not cur_m_df.empty else pd.Series()
    top_custs = cur_cust_sales.nlargest(compare_cust_count).index.tolist() if not cur_cust_sales.empty else []

    cards_data = []
    for cust in top_custs:
        cid_s = cust_id_lookup.get(cust, "")
        cd = build_meeting_comparison_card_data(
            cust, cid_s, cur_m_df, ly_m_df, next_m_df, prev_fy_df,
            tgt_y, tgt_m, next_y, next_m, f_year, max_date
        )
        act_k = f"act_{target_meeting_ym}_{cust}"
        pln_k = f"pln_{target_meeting_ym}_{cust}"
        cd["activity_text"] = st.session_state.get(act_k, "")
        cd["plan_text"] = st.session_state.get(pln_k, "")
        cards_data.append(cd)

    with st.expander("📝 各社の「主な営業活動」「今後の予定」コメントを入力・編集する（帳票＆Excelに反映）"):
        for c_item in cards_data:
            cn = c_item["cust_name"]
            ca1, ca2 = st.columns(2)
            with ca1:
                st.text_area(f"🏢 {cn}：この1ヶ月の主な営業活動", key=f"act_{target_meeting_ym}_{cn}", height=68)
            with ca2:
                st.text_area(f"🏢 {cn}：【今後の予定】", key=f"pln_{target_meeting_ym}_{cn}", height=68)

    for c_item in cards_data:
        cn = c_item["cust_name"]
        c_item["activity_text"] = st.session_state.get(f"act_{target_meeting_ym}_{cn}", "")
        c_item["plan_text"] = st.session_state.get(f"pln_{target_meeting_ym}_{cn}", "")

    st.download_button(
        label=f"📥 上位社 売上比較シート（{target_meeting_ym}）をExcel保存",
        data=generate_meeting_comparison_excel(cards_data, target_meeting_ym),
        file_name=f"営業会議_上位社売上比較シート_{target_meeting_ym}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="dl_btn_meeting_comp_v3"
    )

    st.markdown(f"#### 📊 {tgt_y}年{tgt_m}月度 上位{len(top_custs)}社 売上比較シート（担当: {sel_staff_m}）")
    for idx_c in range(0, len(cards_data), 2):
        cols_card = st.columns(2)
        for sub_i in range(2):
            if idx_c + sub_i < len(cards_data):
                with cols_card[sub_i]:
                    st.markdown(render_meeting_card_html(cards_data[idx_c + sub_i]), unsafe_allow_html=True)
    st.stop()

elif view_type == '📑 営業会議用：上位客先別売上実績':
    if "tgt_ym_topcust_v3" not in st.session_state:
        st.session_state["tgt_ym_topcust_v3"] = ym_start if ym_start in desc_valid_ym else desc_valid_ym[0]

    col_m1, col_m2, col_m3 = st.columns([1.5, 1.5, 1.5])
    with col_m1: target_meeting_ym = st.selectbox("対象月度（期間ショートカット連動）", desc_valid_ym, key="tgt_ym_topcust_v3")
    with col_m2: meeting_cust_count = st.selectbox("表示上位数", [10, 20, 30, 50, 100], index=0, key="meeting_cust_count_v3")
    with col_m3:
        staff_options_topcust = ["（すべて）"] + sorted([x for x in df["担当者"].unique() if x and x != "nan"])
        sel_staff_tc = st.selectbox("担当者指定", staff_options_topcust, key="topcust_staff_sel_v3")

    tgt_y = int(target_meeting_ym[:4]); tgt_m = int(target_meeting_ym[5:7])
    ly_ym = f"{tgt_y - 1}年{tgt_m:02d}月"
    f_year = tgt_y if tgt_m >= 7 else tgt_y - 1
    term_start_ym = f"{f_year}年07月"
    ly_term_start_ym = f"{f_year - 1}年07月"
    ly_term_end_ym = f"{tgt_y - 1}年{tgt_m:02d}月"

    df_scope = df if sel_staff_tc == "（すべて）" else df[df["担当者"] == sel_staff_tc]

    # 🔴 自己参照を完全に排除した安全な抽出
    cur_m_df = df_scope[df_scope["年月"] == target_meeting_ym] if not df_scope.empty else pd.DataFrame()
    ly_m_df = df_scope[df_scope["年月"] == ly_ym] if not df_scope.empty else pd.DataFrame()
    cur_ytd_df = df_scope[(df_scope["年月"] >= term_start_ym) & (df_scope["年月"] <= target_meeting_ym)] if not df_scope.empty else pd.DataFrame()
    ly_ytd_df = df_scope[(df_scope["年月"] >= ly_term_start_ym) & (df_scope["年月"] <= ly_term_end_ym)] if not df_scope.empty else pd.DataFrame()

    cur_cust_sales = cur_m_df.groupby("得意先名")["金額"].sum() if not cur_m_df.empty else pd.Series()
    ly_cust_sales = ly_m_df.groupby("得意先名")["金額"].sum() if not ly_m_df.empty else pd.Series()
    cur_ytd_sales = cur_ytd_df.groupby("得意先名")["金額"].sum() if not cur_ytd_df.empty else pd.Series()
    ly_ytd_sales = ly_ytd_df.groupby("得意先名")["金額"].sum() if not ly_ytd_df.empty else pd.Series()

    target_cust_list = cur_cust_sales.nlargest(meeting_cust_count).index.tolist() if not cur_cust_sales.empty else []
    m_rows = []
    tot_cur_m, tot_ly_m, tot_cur_ytd, tot_ly_ytd = 0, 0, 0, 0

    for rank, cust in enumerate(target_cust_list, 1):
        s_cur = int(cur_cust_sales.get(cust, 0))
        s_ly = int(ly_cust_sales.get(cust, 0))
        ytd_cur = int(cur_ytd_sales.get(cust, 0))
        ytd_ly = int(ly_ytd_sales.get(cust, 0))
        diff = s_cur - s_ly

        tot_cur_m += s_cur
        tot_ly_m += s_ly
        tot_cur_ytd += ytd_cur
        tot_ly_ytd += ytd_ly

        m_rows.append({
            "順位": str(rank),
            "客先名": cust,
            "売上実績": s_cur,
            "前年実績": s_ly,
            "前年比": f"{(s_cur / s_ly * 100):.0f}%" if s_ly != 0 else "-",
            "増減": diff,
            "今期累計実績": ytd_cur,
            "前期累計実績": ytd_ly,
            "累計前年比": f"{(ytd_cur / ytd_ly * 100):.0f}%" if ytd_ly != 0 else "-"
        })

    res_m_df = pd.DataFrame(m_rows)
    tot_row = pd.DataFrame([{
        "順位": "", "客先名": "合計", "売上実績": tot_cur_m, "前年実績": tot_ly_m,
        "前年比": f"{(tot_cur_m / tot_ly_m * 100):.0f}%" if tot_ly_m != 0 else "-", "増減": tot_cur_m - tot_ly_m,
        "今期累計実績": tot_cur_ytd, "前期累計実績": tot_ly_ytd,
        "累計前年比": f"{(tot_cur_ytd / tot_ly_ytd * 100):.0f}%" if tot_ly_ytd != 0 else "-"
    }])
    res_m_df_tot = pd.concat([res_m_df, tot_row], ignore_index=True)
    title_sheet = f"{tgt_m}月度上位{len(target_cust_list)}社客先別売上実績"

    st.download_button(
        label=f"📥 上位客先別売上実績（{target_meeting_ym}）をExcel保存",
        data=generate_meeting_top_cust_excel(res_m_df_tot, title_sheet),
        file_name=f"営業会議_{title_sheet}_{target_meeting_ym}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="dl_btn_top_cust_v3"
    )

    st.markdown(f"### {title_sheet}（担当: {sel_staff_tc}）")
    render_universal_table(res_m_df_tot)
    st.stop()

elif view_type == '📈 商品別 前年実績比較':
    st.markdown(f"#### 📈 商品別 前年実績比較（対象期間: {period_text}）")
    cur_p = filtered.groupby(["商品コード", "商品名"]).agg(当年数量=("数量", "sum"), 当年売上=("金額", "sum"), 当年粗利=("粗利額", "sum")).reset_index() if not filtered.empty else pd.DataFrame(columns=["商品コード", "商品名", "当年数量", "当年売上", "当年粗利"])
    ly_p = ly_filtered.groupby(["商品コード"]).agg(前年数量=("数量", "sum"), 前年売上=("金額", "sum"), 前年粗利=("粗利額", "sum")).reset_index() if not ly_filtered.empty else pd.DataFrame(columns=["商品コード", "前年数量", "前年売上", "前年粗利"])
    comp_p = pd.merge(cur_p, ly_p, on="商品コード", how="outer").fillna(0)
    comp_p["商品名"] = comp_p.apply(lambda r: r["商品名"] if r["商品名"] not in [0, "0", ""] else code_to_name.get(r["商品コード"], r["商品コード"]), axis=1)
    for col_int in ["当年数量", "前年数量", "当年売上", "前年売上", "当年粗利", "前年粗利"]:
        comp_p[col_int] = comp_p[col_int].astype(int)
    comp_p["増減額"] = comp_p["当年売上"] - comp_p["前年売上"]
    comp_p["前年比"] = comp_p.apply(lambda r: f"{(r['当年売上'] / r['前年売上'] * 100):.1f}%" if r["前年売上"] > 0 else "-", axis=1)
    comp_p = comp_p.sort_values(by="当年売上", ascending=False)
    comp_p.insert(0, "順位", range(1, len(comp_p) + 1))
    if limit_val != "全件": comp_p = comp_p.head(int(limit_val))
    show_cp = comp_p[["順位", "商品コード", "商品名", "当年数量", "前年数量", "当年売上", "前年売上", "前年比", "増減額", "当年粗利"]].copy()
    render_universal_table(show_cp)
    st.stop()

elif view_type == '📊 得意先別 前年実績比較':
    st.markdown(f"#### 📊 得意先別 前年実績比較（対象期間: {period_text}）")
    cur_c = filtered.groupby(["得意先名"]).agg(担当者=("担当者", "last"), 当年売上=("金額", "sum"), 当年粗利=("粗利額", "sum"), 当年数量=("数量", "sum")).reset_index() if not filtered.empty else pd.DataFrame(columns=["得意先名", "担当者", "当年売上", "当年粗利", "当年数量"])
    ly_c = ly_filtered.groupby(["得意先名"]).agg(前年売上=("金額", "sum"), 前年粗利=("粗利額", "sum"), 前年数量=("数量", "sum")).reset_index() if not ly_filtered.empty else pd.DataFrame(columns=["得意先名", "前年売上", "前年粗利", "前年数量"])
    comp_c = pd.merge(cur_c, ly_c, on="得意先名", how="outer").fillna(0)
    comp_c["担当者"] = comp_c["担当者"].replace(0, "")
    for col_int in ["当年売上", "前年売上", "当年粗利", "前年粗利", "当年数量", "前年数量"]:
        comp_c[col_int] = comp_c[col_int].astype(int)
    comp_c["増減額"] = comp_c["当年売上"] - comp_c["前年売上"]
    comp_c["前年比"] = comp_c.apply(lambda r: f"{(r['当年売上'] / r['前年売上'] * 100):.1f}%" if r["前年売上"] > 0 else "-", axis=1)
    comp_c = comp_c.sort_values(by="当年売上", ascending=False)
    comp_c.insert(0, "順位", range(1, len(comp_c) + 1))
    if limit_val != "全件": comp_c = comp_c.head(int(limit_val))
    show_cc = comp_c[["順位", "得意先名", "担当者", "当年売上", "前年売上", "前年比", "増減額", "当年粗利", "当年数量"]].copy()
    render_universal_table(show_cc)
    st.stop()

elif view_type == '🧩 得意先×商品 マトリクス':
    st.markdown(f"#### 🧩 得意先×商品 売上マトリクス（対象期間: {period_text}）")
    if not filtered.empty:
        top_c_list = filtered.groupby("得意先名")["金額"].sum().nlargest(30).index.tolist()
        top_p_list = filtered.groupby("商品コード")["金額"].sum().nlargest(20).index.tolist()
        sub_mtx = filtered[filtered["得意先名"].isin(top_c_list) & filtered["商品コード"].isin(top_p_list)]
        piv = sub_mtx.pivot_table(index="得意先名", columns="商品コード", values="金額", aggfunc="sum", fill_value=0)
        piv["合計"] = piv.sum(axis=1)
        piv = piv.sort_values(by="合計", ascending=False).reset_index()
        render_universal_table(piv)
    else:
        st.info("該当期間のデータがありません。")
    st.stop()

elif view_type == '💤 滞留・未販売商品分析（売れてない商品）':
    st.markdown(f"#### 💤 滞留・未販売商品分析（基準日: {max_date.strftime('%Y/%m/%d')}）")
    cutoff_days = st.slider("直近何日間売れていない商品を抽出するか", min_value=30, max_value=365, value=90, step=30)
    cutoff_dt = pd.Timestamp(max_date) - pd.Timedelta(days=cutoff_days)
    dormant_rows = []
    for sku, l_dt in prod_last_seen_dt.items():
        if pd.notna(l_dt) and l_dt < cutoff_dt:
            days_diff = (pd.Timestamp(max_date) - l_dt).days
            dormant_rows.append({
                "商品コード": sku,
                "商品名": code_to_name.get(sku, sku),
                "最終売上日": l_dt.strftime('%Y/%m/%d'),
                "未販売期間": f"{days_diff} 日間",
                "単価原価": int(cost_master_dict.get(sku, 0))
            })
    dorm_df = pd.DataFrame(dormant_rows)
    if not dorm_df.empty:
        dorm_df = dorm_df.sort_values(by="最終売上日", ascending=True).reset_index(drop=True)
        dorm_df.insert(0, "番号", range(1, len(dorm_df) + 1))
        if limit_val != "全件": dorm_df = dorm_df.head(int(limit_val))
    render_universal_table(dorm_df)
    st.stop()

else:
    # 📋 明細一覧
    disp_df = filtered.drop(columns=["年月", "日付_dt", "元商品名"], errors="ignore").sort_values(by="金額", ascending=False) if not filtered.empty else pd.DataFrame()
    if not disp_df.empty:
        if "番号" not in disp_df.columns:
            disp_df.insert(0, "番号", range(1, len(disp_df) + 1))
        if limit_val != "全件": disp_df = disp_df.head(int(limit_val))
    render_universal_table(disp_df)
